"""
MEP Group ERP - Warehouse Inter Sales Chalan Report Collector
Sequential Date-Wise Extraction & Premium Excel Generator
Author: Antigravity
"""

import os
import re
import json
import time
from datetime import datetime, timedelta
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_URL = "https://www.mepgrouperp.com/1027"
LOGIN_URL = f"{BASE_URL}/login/pages/main/index.php"
REPORT_LIST_URL = f"{BASE_URL}/warehouse_mod/pages/inter_sales/report_list.php"
MASTER_REPORT_URL = f"{BASE_URL}/warehouse_mod/pages/inter_sales/master_report.php"
INVOICE_URL = f"{BASE_URL}/warehouse_mod/pages/inter_sales/invoice.php?req_no="
PRINT_VIEW_URL = f"{BASE_URL}/warehouse_mod/pages/inter_sales/print_view.php?req_no="

from erp_credentials import get_erp_credentials

DEFAULT_CREDENTIALS = get_erp_credentials()

CACHE_FILE = "inter_sales_chalan_data.json"
EXCEL_OUTPUT_FILE = "Inter_Sales_Chalan_Report_Sep2026.xlsx"

# Visual Styles for Premium Excel
FONT_NAME = "Segoe UI"
COLOR_NAVY = "1B365D"       # Corporate Dark Navy
COLOR_SECONDARY = "2B5B84"  # Medium Slate Navy
COLOR_ACCENT = "008080"     # Teal / Accent
COLOR_LIGHT_BG = "F4F7FB"   # Soft Blue-Gray
COLOR_ZEBRA = "F9FBFC"      # Subtle Alternating Row
COLOR_BORDER = "D1D5DB"     # Soft Gray
COLOR_ACTIVE_BG = "E6F4EA"  # Soft Green
COLOR_ACTIVE_FG = "137333"  # Dark Green
COLOR_NODATA_BG = "FEE2E2"  # Soft Red
COLOR_NODATA_FG = "B91C1C"  # Dark Red
COLOR_DUP_BG = "FEF3C7"     # Amber / Yellow
COLOR_DUP_FG = "B45309"     # Dark Amber

thin_border = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_BORDER),
    bottom=Side(style='thin', color=COLOR_BORDER)
)

header_border = Border(
    left=Side(style='thin', color=COLOR_NAVY),
    right=Side(style='thin', color=COLOR_NAVY),
    top=Side(style='medium', color=COLOR_NAVY),
    bottom=Side(style='medium', color=COLOR_NAVY)
)


class InterSalesChalanCollector:
    def __init__(self, credentials=None):
        self.creds = credentials or DEFAULT_CREDENTIALS
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
        })
        self.daily_reports = []   # Date -> Company -> Destination -> Chalans
        self.flat_chalans = []    # Flattened chalan records
        self.itemized_lines = []  # Detailed invoice items
        self.seen_chalans = {}    # Duplicate tracker

    def login(self):
        """Authenticates with MEP ERP."""
        print("[*] Authenticating with MEP Group ERP...")
        res = self.session.post(LOGIN_URL, data=self.creds, allow_redirects=True, timeout=20)
        # Verify access
        test = self.session.get(REPORT_LIST_URL, timeout=15)
        if "InterSales Chalan Report" not in test.text:
            raise Exception("Failed to access Inter Sales Chalan Report. Please check credentials.")
        print("[+] Logged in successfully and verified access to Inter Sales Chalan Report!")
        return True

    def collect_date_range(self, start_date_str="2026-09-01", end_date_str="2026-09-13",
                           company="Printing and Packaging",
                           target_destinations=None,
                           fetch_item_details=True):
        """
        Executes strictly sequential date-wise report collection.
        Processes one date at a time from start_date up to end_date.
        """
        if target_destinations is None:
            target_destinations = ["MEP", "Fan", "MEP Light"]

        start_dt = datetime.strptime(start_date_str, "%Y-%m-%d")
        end_dt = datetime.strptime(end_date_str, "%Y-%m-%d")

        print(f"\n[*] Starting Sequential Date-Wise Collection:")
        print(f"    Date Range: {start_date_str} to {end_date_str} ({(end_dt - start_dt).days + 1} days)")
        print(f"    Company Filter: {company}")
        print(f"    Target Destinations: {target_destinations}\n")

        self.daily_reports = []
        self.flat_chalans = []
        self.seen_chalans = {}

        curr = start_dt
        while curr <= end_dt:
            date_str = curr.strftime("%Y-%m-%d")
            day_name = curr.strftime("%A")
            print(f"[*] Processing Date: {date_str} ({day_name})...")

            # Form parameters: report=201, group_for=1 (Printing and Packaging)
            post_data = {
                "report": "201",
                "group_for": "1",  # 1 corresponds to 'Printing and Packaging'
                "item_id": "",
                "f_date": date_str,
                "t_date": date_str,
                "warehouse_id": "",
                "submit": "Report"
            }

            try:
                res = self.session.post(MASTER_REPORT_URL, data=post_data, timeout=30)
                day_data = self._parse_day_report(res.text, date_str, company, target_destinations)
                self.daily_reports.append(day_data)

                total_day_chalans = sum(len(dest_info["chalans"]) for dest_info in day_data["destinations"].values())
                print(f"    -> Extracted {total_day_chalans} chalans for {date_str} across {len(day_data['destinations'])} destinations.")
            except Exception as e:
                print(f"[!] Error on {date_str}: {e}")
                self.daily_reports.append({
                    "date": date_str,
                    "day_name": day_name,
                    "company": company,
                    "status": f"Error: {str(e)}",
                    "destinations": {}
                })

            curr += timedelta(days=1)

        print(f"\n[+] Completed date-wise extraction: {len(self.flat_chalans)} total chalan rows extracted.")

        # Concurrently fetch invoice line items if requested
        if fetch_item_details and self.flat_chalans:
            self._fetch_all_invoice_items()

        return self.daily_reports

    def _normalize_dest(self, dest_str):
        """Normalizes destination names (e.g. FAN -> Fan, Mep Light -> MEP Light)."""
        d = (dest_str or "").strip()
        if d.upper() == "MEP":
            return "MEP"
        elif d.upper() == "FAN":
            return "Fan"
        elif "LIGHT" in d.upper():
            return "MEP Light"
        return d

    def _parse_day_report(self, html, date_str, company, target_destinations):
        """Parses the panel report table for a single date."""
        soup = BeautifulSoup(html, "html.parser")
        table = soup.find("table")
        day_name = datetime.strptime(date_str, "%Y-%m-%d").strftime("%A")

        day_data = {
            "date": date_str,
            "day_name": day_name,
            "company": company,
            "status": "Success",
            "destinations": {}
        }

        # Initialize target destinations
        for t in target_destinations:
            norm_t = self._normalize_dest(t)
            day_data["destinations"][norm_t] = {
                "destination_name": norm_t,
                "status": "No Data",
                "chalans": []
            }

        if not table:
            day_data["status"] = "No Data"
            return day_data

        rows = table.find_all("tr")
        # Row 0: Title, Row 1: Headers ('S/L', 'Date', 'InterSales No', 'From', 'To', 'Sales_Amount')
        extracted_rows = []
        for tr in rows[2:]:
            tds = tr.find_all(["td", "th"])
            cells = [c.get_text(strip=True) for c in tds]
            if not cells or cells[0] == "Total":
                continue

            sl = cells[0]
            r_date = cells[1]

            # InterSales No & Links
            chalan_td = tds[2] if len(tds) > 2 else None
            chalan_no = ""
            chalan_url = ""
            invoice_url = ""
            if chalan_td:
                for a in chalan_td.find_all("a"):
                    txt = a.get_text(strip=True)
                    href = a.get("href", "")
                    if "print_view.php" in href:
                        chalan_no = txt
                        chalan_url = f"{BASE_URL}/warehouse_mod/pages/inter_sales/{href}"
                    elif "invoice.php" in href:
                        invoice_url = f"{BASE_URL}/warehouse_mod/pages/inter_sales/{href}"
                if not chalan_no:
                    m = re.search(r"\d+", chalan_td.get_text(strip=True))
                    chalan_no = m.group(0) if m else chalan_td.get_text(strip=True)

            company_from = cells[3] if len(cells) > 3 else company
            dest_raw = cells[4] if len(cells) > 4 else ""
            dest_norm = self._normalize_dest(dest_raw)

            sales_amt_str = cells[5] if len(cells) > 5 else "0.00"
            try:
                sales_amt = float(sales_amt_str.replace(",", ""))
            except ValueError:
                sales_amt = 0.0

            # Duplicate Detection
            is_dup = False
            dup_status = "Unique"
            if chalan_no in self.seen_chalans:
                is_dup = True
                first_date, first_sl = self.seen_chalans[chalan_no]
                dup_status = f"DUPLICATE (First seen: {first_date}, S/L {first_sl})"
            else:
                self.seen_chalans[chalan_no] = (r_date, sl)

            record = {
                "date": r_date,
                "day_name": day_name,
                "company": company_from,
                "destination_raw": dest_raw,
                "destination": dest_norm,
                "original_sl": int(sl) if sl.isdigit() else sl,
                "chalan_no": chalan_no,
                "sales_amount": sales_amt,
                "sales_amount_str": sales_amt_str,
                "is_duplicate": is_dup,
                "duplicate_status": dup_status,
                "data_status": "Data Available",
                "chalan_url": chalan_url,
                "invoice_url": invoice_url,
                "items": []
            }
            extracted_rows.append(record)
            self.flat_chalans.append(record)

            # Assign to day_data destination
            if dest_norm not in day_data["destinations"]:
                day_data["destinations"][dest_norm] = {
                    "destination_name": dest_norm,
                    "status": "Data Available",
                    "chalans": []
                }
            day_data["destinations"][dest_norm]["status"] = "Data Available"
            day_data["destinations"][dest_norm]["chalans"].append(record)

        if not extracted_rows:
            day_data["status"] = "No Data"

        return day_data

    def _fetch_all_invoice_items(self, max_workers=10):
        """Fetches line-item breakdown from invoice.php for all extracted chalans."""
        print(f"[*] Fetching invoice line items for {len(self.flat_chalans)} chalans using {max_workers} threads...")
        cookies = self.session.cookies.get_dict()
        chalan_map = {c["chalan_no"]: c for c in self.flat_chalans}

        def worker(chalan_rec):
            c_no = chalan_rec["chalan_no"]
            if not c_no:
                return c_no, []
            s = requests.Session()
            s.cookies.update(cookies)
            u = f"{INVOICE_URL}{c_no}"
            try:
                r = s.get(u, timeout=15)
                soup = BeautifulSoup(r.text, "html.parser")
                items = []
                for tbl in soup.find_all("table"):
                    rows = tbl.find_all("tr")
                    headers = [c.get_text(strip=True) for c in rows[0].find_all(["td", "th"])] if rows else []
                    if "Item Code" in headers and "Product Name" in headers:
                        for tr in rows[1:]:
                            tds = [c.get_text(strip=True) for c in tr.find_all(["td", "th"])]
                            if tds and tds[0] != "Total:" and len(tds) >= 7:
                                try:
                                    qty_f = float(tds[4].replace(",", ""))
                                except ValueError:
                                    qty_f = tds[4]
                                try:
                                    rate_f = float(tds[5].replace(",", ""))
                                except ValueError:
                                    rate_f = tds[5]
                                try:
                                    tot_f = float(tds[6].replace(",", ""))
                                except ValueError:
                                    tot_f = tds[6]

                                items.append({
                                    "sl": tds[0],
                                    "item_code": tds[1],
                                    "product_name": tds[2],
                                    "unit": tds[3],
                                    "qty": qty_f,
                                    "rate": rate_f,
                                    "total_amt": tot_f,
                                    "vat_pct": tds[7] if len(tds) > 7 else "",
                                    "vat_amt": tds[8] if len(tds) > 8 else ""
                                })
                return c_no, items
            except Exception as err:
                return c_no, []

        completed = 0
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_c = {executor.submit(worker, c): c for c in self.flat_chalans}
            for future in as_completed(future_to_c):
                c_no, items = future.result()
                if c_no in chalan_map:
                    chalan_map[c_no]["items"] = items
                    for it in items:
                        self.itemized_lines.append({
                            "date": chalan_map[c_no]["date"],
                            "company": chalan_map[c_no]["company"],
                            "destination": chalan_map[c_no]["destination"],
                            "chalan_no": c_no,
                            **it
                        })
                completed += 1

        print(f"[+] Successfully fetched {len(self.itemized_lines)} line items across {len(self.flat_chalans)} chalans.")

    def save_cache(self, filepath=CACHE_FILE):
        """Saves scraped report data to JSON cache."""
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump({
                "collected_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "total_chalans": len(self.flat_chalans),
                "total_amount": sum(c["sales_amount"] for c in self.flat_chalans),
                "daily_reports": self.daily_reports,
                "flat_chalans": self.flat_chalans,
                "itemized_lines": self.itemized_lines
            }, f, indent=2, ensure_ascii=False)
        print(f"[+] Saved cache to {filepath}")

    def load_cache(self, filepath=CACHE_FILE):
        """Loads data from JSON cache."""
        if os.path.exists(filepath):
            with open(filepath, "r", encoding="utf-8") as f:
                d = json.load(f)
                self.daily_reports = d.get("daily_reports", [])
                self.flat_chalans = d.get("flat_chalans", [])
                self.itemized_lines = d.get("itemized_lines", [])
                return True
        return False

    def export_to_premium_excel(self, output_file=EXCEL_OUTPUT_FILE):
        """Builds executive-grade multi-sheet Excel workbook."""
        print(f"[*] Generating Premium Excel Workbook: {output_file}...")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # 1. Executive Summary & KPIs
        ws_dash = wb.create_sheet(title="Executive Summary")
        ws_dash.views.sheetView[0].showGridLines = True
        self._build_dashboard_sheet(ws_dash)

        # 2. Panel Report (Date-Wise)
        ws_panel = wb.create_sheet(title="Panel Report (Date-Wise)")
        ws_panel.views.sheetView[0].showGridLines = True
        self._build_panel_report_sheet(ws_panel)

        # 3. Date & Destination Matrix
        ws_matrix = wb.create_sheet(title="Date-Destination Matrix")
        ws_matrix.views.sheetView[0].showGridLines = True
        self._build_matrix_sheet(ws_matrix)

        # 4. Chalan Itemized Breakdown
        ws_items = wb.create_sheet(title="Chalan Itemized Breakdown")
        ws_items.views.sheetView[0].showGridLines = True
        self._build_itemized_sheet(ws_items)

        # 5. Audit & No Data Log
        ws_audit = wb.create_sheet(title="Audit & No Data Log")
        ws_audit.views.sheetView[0].showGridLines = True
        self._build_audit_sheet(ws_audit)

        wb.save(output_file)
        print(f"[+] Successfully exported {output_file}!")
        return output_file

    def _build_dashboard_sheet(self, ws):
        """Executive KPI & Daily Performance Calendar Summary."""
        # Title Banner
        ws.merge_cells("A1:H2")
        b = ws["A1"]
        b.value = "MEP GROUP  |  Inter Sales Chalan Executive Performance Report"
        b.font = Font(name=FONT_NAME, size=16, bold=True, color="FFFFFF")
        b.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        b.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws["A3"].value = f"Source: Warehouse Module → Inter Sales  |  Company: Printing and Packaging  |  Generated: {datetime.now().strftime('%d-%b-%Y %I:%M %p')}"
        ws["A3"].font = Font(name=FONT_NAME, size=9, italic=True, color="6B7280")

        total_chalans = len(self.flat_chalans)
        total_amount = sum(c["sales_amount"] for c in self.flat_chalans)
        days_count = len(self.daily_reports)
        active_days = sum(1 for d in self.daily_reports if d.get("status") == "Success")
        no_data_days = days_count - active_days

        # Destination Stats
        dest_stats = {}
        for c in self.flat_chalans:
            d = c["destination"]
            if d not in dest_stats:
                dest_stats[d] = {"count": 0, "amount": 0.0}
            dest_stats[d]["count"] += 1
            dest_stats[d]["amount"] += c["sales_amount"]

        # KPI Metric Cards
        kpis = [
            ("Total Chalans", total_chalans, f"{days_count} Days Evaluated", "1B365D"),
            ("Total Sales Value", f"BDT {total_amount:,.2f}", "Printing & Packaging", "137333"),
            ("Active Days with Data", active_days, f"{no_data_days} Holidays / No Data Days", "2B5B84"),
            ("Total Item Lines", len(self.itemized_lines), "Product items dispatched", "008080"),
        ]

        col_pairs = [("B", "C"), ("D", "E"), ("F", "G"), ("H", "I")]
        for idx, (label, val, sub, color_hex) in enumerate(kpis):
            c1, c2 = col_pairs[idx]
            ws.merge_cells(f"{c1}5:{c2}5")
            ws.merge_cells(f"{c1}6:{c2}6")
            ws.merge_cells(f"{c1}7:{c2}7")

            card_val = ws[f"{c1}5"]
            card_val.value = val
            card_val.font = Font(name=FONT_NAME, size=18, bold=True, color=color_hex)
            card_val.alignment = Alignment(horizontal="center", vertical="center")

            card_lbl = ws[f"{c1}6"]
            card_lbl.value = label
            card_lbl.font = Font(name=FONT_NAME, size=10.5, bold=True, color="374151")
            card_lbl.alignment = Alignment(horizontal="center", vertical="center")

            card_sub = ws[f"{c1}7"]
            card_sub.value = sub
            card_sub.font = Font(name=FONT_NAME, size=8, color="6B7280")
            card_sub.alignment = Alignment(horizontal="center", vertical="center")

            for r in range(5, 8):
                for col in [c1, c2]:
                    ws[f"{col}{r}"].fill = PatternFill(start_color=COLOR_LIGHT_BG, end_color=COLOR_LIGHT_BG, fill_type="solid")
                    ws[f"{col}{r}"].border = thin_border

        # Destination Performance Breakdown Table
        ws["B9"].value = "DESTINATION PERFORMANCE BREAKDOWN"
        ws["B9"].font = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_NAVY)

        dest_headers = ["To (Destination)", "Total Chalans", "% Share (Volume)", "Total Sales Amount (BDT)", "% Share (Value)"]
        for c_idx, h in enumerate(dest_headers, start=2):
            cell = ws.cell(row=10, column=c_idx, value=h)
            cell.font = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            cell.alignment = Alignment(horizontal="center" if c_idx in [3, 5] else ("right" if c_idx == 4 else "left"), vertical="center")
            cell.border = header_border

        curr_r = 11
        for dest_name, st in sorted(dest_stats.items(), key=lambda x: x[1]["amount"], reverse=True):
            bg = COLOR_ZEBRA if curr_r % 2 == 1 else "FFFFFF"
            pct_vol = (st["count"] / total_chalans * 100) if total_chalans else 0
            pct_val = (st["amount"] / total_amount * 100) if total_amount else 0

            row_data = [dest_name, st["count"], f"{pct_vol:.1f}%", st["amount"], f"{pct_val:.1f}%"]
            for c_idx, v in enumerate(row_data, start=2):
                cell = ws.cell(row=curr_r, column=c_idx, value=v)
                cell.font = Font(name=FONT_NAME, size=9.5)
                cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center" if c_idx in [3, 5] else ("right" if c_idx in [2, 4] else "left"), vertical="center")
                if c_idx == 4:
                    cell.number_format = "#,##0.00"
            curr_r += 1

        # Total Row
        ws.cell(row=curr_r, column=2, value="Grand Total").font = Font(name=FONT_NAME, size=9.5, bold=True)
        ws.cell(row=curr_r, column=3, value=total_chalans).font = Font(name=FONT_NAME, size=9.5, bold=True)
        ws.cell(row=curr_r, column=4, value="100.0%").font = Font(name=FONT_NAME, size=9.5, bold=True)
        tot_cell = ws.cell(row=curr_r, column=5, value=total_amount)
        tot_cell.font = Font(name=FONT_NAME, size=9.5, bold=True)
        tot_cell.number_format = "#,##0.00"
        ws.cell(row=curr_r, column=6, value="100.0%").font = Font(name=FONT_NAME, size=9.5, bold=True)

        for c_idx in range(2, 7):
            c = ws.cell(row=curr_r, column=c_idx)
            c.fill = PatternFill(start_color=COLOR_LIGHT_BG, end_color=COLOR_LIGHT_BG, fill_type="solid")
            c.border = Border(top=Side(style='thin'), bottom=Side(style='double'))
            if c_idx in [3, 4, 6]:
                c.alignment = Alignment(horizontal="center" if c_idx != 5 else "right")

        # Daily Calendar Matrix
        curr_r += 3
        ws.cell(row=curr_r, column=2, value="DAILY CALENDAR DISPATCH LOG (01-SEP TO 13-SEP)").font = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_NAVY)
        curr_r += 1

        cal_headers = ["Date", "Day of Week", "MEP Chalans", "Fan Chalans", "MEP Light Chalans", "Other Chalans", "Day Total Chalans", "Day Total (BDT)", "Status"]
        for c_idx, h in enumerate(cal_headers, start=2):
            c = ws.cell(row=curr_r, column=c_idx, value=h)
            c.font = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color=COLOR_SECONDARY, end_color=COLOR_SECONDARY, fill_type="solid")
            c.alignment = Alignment(horizontal="center" if c_idx != 9 else "right", vertical="center")
            c.border = header_border

        curr_r += 1
        for d in self.daily_reports:
            d_str = d["date"]
            d_name = d["day_name"]
            dests = d.get("destinations", {})

            mep_count = len(dests.get("MEP", {}).get("chalans", []))
            fan_count = len(dests.get("Fan", {}).get("chalans", []))
            light_count = len(dests.get("MEP Light", {}).get("chalans", []))
            other_count = sum(len(dests[k].get("chalans", [])) for k in dests if k not in ["MEP", "Fan", "MEP Light"])
            day_total_chalans = mep_count + fan_count + light_count + other_count
            day_total_amt = sum(c["sales_amount"] for dest_info in dests.values() for c in dest_info.get("chalans", []))
            status = "No Data" if day_total_chalans == 0 else "Active"

            bg = COLOR_ZEBRA if curr_r % 2 == 1 else "FFFFFF"
            vals = [d_str, d_name, mep_count, fan_count, light_count, other_count, day_total_chalans, day_total_amt, status]

            for c_idx, val in enumerate(vals, start=2):
                c = ws.cell(row=curr_r, column=c_idx, value=val)
                c.font = Font(name=FONT_NAME, size=9.5)
                c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                c.border = thin_border
                c.alignment = Alignment(horizontal="center" if c_idx in [2, 3, 4, 5, 6, 7, 8, 10] else ("right" if c_idx == 9 else "left"), vertical="center")

                if c_idx == 9:
                    c.number_format = "#,##0.00"
                if c_idx == 10:
                    if status == "Active":
                        c.fill = PatternFill(start_color=COLOR_ACTIVE_BG, end_color=COLOR_ACTIVE_BG, fill_type="solid")
                        c.font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_ACTIVE_FG)
                    else:
                        c.fill = PatternFill(start_color=COLOR_NODATA_BG, end_color=COLOR_NODATA_BG, fill_type="solid")
                        c.font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_NODATA_FG)

            curr_r += 1

        self._auto_adjust_columns(ws, min_width=14, max_width=40)

    def _build_panel_report_sheet(self, ws):
        """Hierarchical table: Date -> Company -> To -> S/L -> Chalan No -> Sales Amount."""
        ws.merge_cells("A1:K1")
        title = ws["A1"]
        title.value = "MEP GROUP  |  Inter Sales Chalan Report (Date → Company → To → Chalan)"
        title.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
        title.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 30

        headers = [
            ("Date", 13, "center"),
            ("Day", 11, "center"),
            ("Company (From)", 24, "left"),
            ("To (Destination)", 18, "left"),
            ("Panel S/L", 10, "center"),
            ("Chalan / InterSales No", 22, "center"),
            ("Sales Amount (BDT)", 18, "right"),
            ("Duplicate Status", 16, "center"),
            ("Data Status", 14, "center"),
            ("Chalan Link", 15, "center"),
            ("Invoice Link", 15, "center"),
        ]

        ws.row_dimensions[3].height = 24
        for c_idx, (h_title, _, align) in enumerate(headers, start=1):
            c = ws.cell(row=3, column=c_idx, value=h_title)
            c.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            c.alignment = Alignment(horizontal=align, vertical="center")
            c.border = header_border

        curr_r = 4
        # Iterate in strict sequence: Date -> Company -> Destination -> Chalans
        for d in self.daily_reports:
            d_str = d["date"]
            d_name = d["day_name"]
            comp = d["company"]
            dests = d.get("destinations", {})

            # For each required destination
            target_keys = ["MEP", "Fan", "MEP Light"]
            # Also append any other destinations that exist on that date
            all_keys = target_keys + [k for k in dests.keys() if k not in target_keys]

            for dest_key in all_keys:
                dest_info = dests.get(dest_key, {"destination_name": dest_key, "status": "No Data", "chalans": []})
                chalans = dest_info.get("chalans", [])

                if not chalans:
                    # Record "No Data" row for this date and destination
                    bg = COLOR_NODATA_BG
                    vals = [d_str, d_name, comp, dest_key, "-", "-", 0.0, "Unique", "No Data", "-", "-"]
                    ws.row_dimensions[curr_r].height = 19
                    for c_idx, v in enumerate(vals, start=1):
                        c = ws.cell(row=curr_r, column=c_idx, value=v)
                        c.font = Font(name=FONT_NAME, size=9.5, italic=(c_idx in [5, 6, 9]))
                        c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                        c.border = thin_border
                        c.alignment = Alignment(horizontal=headers[c_idx-1][2], vertical="center")
                        if c_idx == 7:
                            c.number_format = "#,##0.00"
                        if c_idx == 9:
                            c.font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_NODATA_FG)
                    curr_r += 1
                else:
                    for ch in chalans:
                        bg = COLOR_ZEBRA if curr_r % 2 == 1 else "FFFFFF"
                        if ch.get("is_duplicate"):
                            bg = COLOR_DUP_BG

                        vals = [
                            d_str,
                            d_name,
                            comp,
                            dest_key,
                            ch.get("original_sl", ""),
                            ch.get("chalan_no", ""),
                            ch.get("sales_amount", 0.0),
                            ch.get("duplicate_status", "Unique"),
                            ch.get("data_status", "Data Available"),
                            ch.get("chalan_url", ""),
                            ch.get("invoice_url", "")
                        ]

                        ws.row_dimensions[curr_r].height = 19
                        for c_idx, v in enumerate(vals, start=1):
                            c = ws.cell(row=curr_r, column=c_idx, value=v)
                            c.font = Font(name=FONT_NAME, size=9.5)
                            c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                            c.border = thin_border
                            c.alignment = Alignment(horizontal=headers[c_idx-1][2], vertical="center")

                            if c_idx == 6:
                                c.font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_NAVY)
                            if c_idx == 7:
                                c.number_format = "#,##0.00"
                            if c_idx == 8 and ch.get("is_duplicate"):
                                c.font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_DUP_FG)
                            if c_idx == 9:
                                c.font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_ACTIVE_FG)
                            if c_idx in [10, 11] and v and v != "-":
                                c.hyperlink = v
                                c.value = "View " + ("Chalan" if c_idx == 10 else "Invoice")
                                c.font = Font(name=FONT_NAME, size=9, color="0000FF", underline="single")

                        curr_r += 1

        ws.auto_filter.ref = f"A3:K{curr_r - 1}"
        ws.freeze_panes = "A4"

        for c_idx, (_, width, _) in enumerate(headers, start=1):
            col_letter = get_column_letter(c_idx)
            ws.column_dimensions[col_letter].width = width

    def _build_matrix_sheet(self, ws):
        """Cross-tabulation Matrix: Date vs Destination."""
        ws.merge_cells("A1:K1")
        title = ws["A1"]
        title.value = "MEP GROUP  |  Date vs Destination Revenue & Chalan Matrix"
        title.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
        title.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 30

        # Subheaders: MEP (Qty, BDT), Fan (Qty, BDT), MEP Light (Qty, BDT), Other (Qty, BDT), Total (Qty, BDT)
        ws.merge_cells("A3:A4")
        ws["A3"] = "Date"
        ws.merge_cells("B3:C3")
        ws["B3"] = "MEP"
        ws.merge_cells("D3:E3")
        ws["D3"] = "Fan"
        ws.merge_cells("F3:G3")
        ws["F3"] = "MEP Light"
        ws.merge_cells("H3:I3")
        ws["H3"] = "Other Destinations"
        ws.merge_cells("J3:K3")
        ws["J3"] = "Total Dispatches"

        sub_cols = ["Chalans", "Amount (BDT)"] * 5
        for i, sc in enumerate(sub_cols, start=2):
            ws.cell(row=4, column=i, value=sc)

        for r in [3, 4]:
            for col in range(1, 12):
                c = ws.cell(row=r, column=col)
                c.font = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
                c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border = header_border

        curr_r = 5
        for d in self.daily_reports:
            d_str = d["date"]
            dests = d.get("destinations", {})

            mep_chalans = dests.get("MEP", {}).get("chalans", [])
            mep_amt = sum(c["sales_amount"] for c in mep_chalans)

            fan_chalans = dests.get("Fan", {}).get("chalans", [])
            fan_amt = sum(c["sales_amount"] for c in fan_chalans)

            light_chalans = dests.get("MEP Light", {}).get("chalans", [])
            light_amt = sum(c["sales_amount"] for c in light_chalans)

            other_chalans = [c for k, v in dests.items() if k not in ["MEP", "Fan", "MEP Light"] for c in v.get("chalans", [])]
            other_amt = sum(c["sales_amount"] for c in other_chalans)

            tot_chalans = len(mep_chalans) + len(fan_chalans) + len(light_chalans) + len(other_chalans)
            tot_amt = mep_amt + fan_amt + light_amt + other_amt

            bg = COLOR_ZEBRA if curr_r % 2 == 1 else "FFFFFF"
            row_vals = [
                d_str,
                len(mep_chalans), mep_amt,
                len(fan_chalans), fan_amt,
                len(light_chalans), light_amt,
                len(other_chalans), other_amt,
                tot_chalans, tot_amt
            ]

            for c_idx, v in enumerate(row_vals, start=1):
                c = ws.cell(row=curr_r, column=c_idx, value=v)
                c.font = Font(name=FONT_NAME, size=9.5)
                c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                c.border = thin_border
                c.alignment = Alignment(horizontal="center" if c_idx in [1, 2, 4, 6, 8, 10] else "right", vertical="center")
                if c_idx in [3, 5, 7, 9, 11]:
                    c.number_format = "#,##0.00"

            curr_r += 1

        self._auto_adjust_columns(ws, min_width=12, max_width=25)

    def _build_itemized_sheet(self, ws):
        """Line items extracted from invoice.php."""
        ws.merge_cells("A1:M1")
        title = ws["A1"]
        title.value = "MEP GROUP  |  Chalan Itemized Breakdown (All Product Lines Dispatched)"
        title.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
        title.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 30

        headers = [
            ("Date", 13, "center"),
            ("Company", 22, "left"),
            ("Destination", 16, "left"),
            ("Chalan No", 14, "center"),
            ("Item SL", 8, "center"),
            ("Item Code", 16, "center"),
            ("Product Description", 38, "left"),
            ("Unit", 8, "center"),
            ("Quantity", 14, "right"),
            ("Unit Rate (BDT)", 15, "right"),
            ("Total Amount (BDT)", 18, "right"),
            ("VAT %", 9, "center"),
            ("VAT Amount (BDT)", 16, "right"),
        ]

        ws.row_dimensions[3].height = 24
        for c_idx, (h_title, _, align) in enumerate(headers, start=1):
            c = ws.cell(row=3, column=c_idx, value=h_title)
            c.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            c.alignment = Alignment(horizontal=align, vertical="center")
            c.border = header_border

        curr_r = 4
        for item in self.itemized_lines:
            bg = COLOR_ZEBRA if curr_r % 2 == 1 else "FFFFFF"
            vals = [
                item.get("date", ""),
                item.get("company", ""),
                item.get("destination", ""),
                item.get("chalan_no", ""),
                item.get("sl", 1),
                item.get("item_code", ""),
                item.get("product_name", ""),
                item.get("unit", ""),
                item.get("qty", 0.0),
                item.get("rate", 0.0),
                item.get("total_amt", 0.0),
                item.get("vat_pct", ""),
                item.get("vat_amt", "")
            ]

            ws.row_dimensions[curr_r].height = 19
            for c_idx, v in enumerate(vals, start=1):
                c = ws.cell(row=curr_r, column=c_idx, value=v)
                c.font = Font(name=FONT_NAME, size=9.5)
                c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                c.border = thin_border
                c.alignment = Alignment(horizontal=headers[c_idx-1][2], vertical="center")

                if c_idx in [9, 10, 11]:
                    c.number_format = "#,##0.00"

            curr_r += 1

        if curr_r == 4:
            ws.cell(row=4, column=1, value="No itemized invoice lines found.").font = Font(name=FONT_NAME, italic=True)
        else:
            ws.auto_filter.ref = f"A3:M{curr_r - 1}"
            ws.freeze_panes = "A4"

        for c_idx, (_, width, _) in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(c_idx)].width = width

    def _build_audit_sheet(self, ws):
        """Audit log of dates, holidays, duplicates, and verification status."""
        ws.merge_cells("A1:G1")
        title = ws["A1"]
        title.value = "MEP GROUP  |  Audit Trail & No-Data Log (Data Integrity Verification)"
        title.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
        title.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 30

        headers = [
            ("Date", 14, "center"),
            ("Day of Week", 14, "center"),
            ("Status", 16, "center"),
            ("Total Chalans Found", 18, "center"),
            ("Duplicate Count", 16, "center"),
            ("Destination Distribution", 36, "left"),
            ("Audit Remarks", 32, "left"),
        ]

        ws.row_dimensions[3].height = 24
        for c_idx, (h_title, _, align) in enumerate(headers, start=1):
            c = ws.cell(row=3, column=c_idx, value=h_title)
            c.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            c.alignment = Alignment(horizontal=align, vertical="center")
            c.border = header_border

        curr_r = 4
        for d in self.daily_reports:
            d_str = d["date"]
            d_name = d["day_name"]
            dests = d.get("destinations", {})
            total_ch = sum(len(dest_info.get("chalans", [])) for dest_info in dests.values())
            dup_count = sum(1 for dest_info in dests.values() for c in dest_info.get("chalans", []) if c.get("is_duplicate"))

            dist_parts = [f"{k}: {len(v.get('chalans', []))}" for k, v in dests.items() if len(v.get('chalans', [])) > 0]
            dist_str = ", ".join(dist_parts) if dist_parts else "None"

            if total_ch == 0:
                status = "No Data"
                remarks = "Weekly Holiday (Friday) or Non-Dispatch Day" if d_name == "Friday" else "No inter-sales dispatches recorded"
                bg = COLOR_NODATA_BG
            else:
                status = "Verified Active"
                remarks = f"Complete data extracted ({total_ch} chalans, {dup_count} duplicates)"
                bg = COLOR_ACTIVE_BG

            vals = [d_str, d_name, status, total_ch, dup_count, dist_str, remarks]
            for c_idx, v in enumerate(vals, start=1):
                c = ws.cell(row=curr_r, column=c_idx, value=v)
                c.font = Font(name=FONT_NAME, size=9.5)
                c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                c.border = thin_border
                c.alignment = Alignment(horizontal=headers[c_idx-1][2], vertical="center")
                if c_idx == 3:
                    c.font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_ACTIVE_FG if total_ch > 0 else COLOR_NODATA_FG)

            curr_r += 1

        ws.auto_filter.ref = f"A3:G{curr_r - 1}"
        ws.freeze_panes = "A4"

        for c_idx, (_, width, _) in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(c_idx)].width = width

    def _auto_adjust_columns(self, ws, min_width=10, max_width=45):
        """Auto-fits column widths based on cell content."""
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2] and cell.coordinate != "A1":
                    continue
                if cell.value:
                    val_str = str(cell.value)
                    if len(val_str) > max_len and len(val_str) < 80:
                        max_len = len(val_str)
            adjusted_width = max(max_len + 3, min_width)
            ws.column_dimensions[col_letter].width = min(adjusted_width, max_width)


def run_inter_sales_extraction(start_date="2026-09-01", end_date="2026-09-13"):
    collector = InterSalesChalanCollector()
    collector.login()
    collector.collect_date_range(start_date_str=start_date, end_date_str=end_date)
    collector.save_cache()
    collector.export_to_premium_excel()
    return collector


if __name__ == "__main__":
    t0 = time.time()
    run_inter_sales_extraction()
    print(f"\n[DONE] Inter Sales Chalan extraction & Premium Excel generated in {time.time()-t0:.2f}s!")
