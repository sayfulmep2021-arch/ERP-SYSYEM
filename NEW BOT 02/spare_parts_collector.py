"""
MEP Group ERP - Spare Parts Stock Summary Report with Rate (3224)
Workflow: Login -> Warehouse -> Report -> Warehouse Reports -> Spare parts Stock Summary Report with rate (3224) -> Company: FAN
Date Range: 1st of Current Month to Running Date (Today)
Author: Antigravity
"""

import os
import json
import time
from datetime import datetime
import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_URL = "https://www.mepgrouperp.com/1027"
LOGIN_URL = f"{BASE_URL}/login/pages/main/index.php"
REPORT_URL = f"{BASE_URL}/warehouse_mod/pages/report/master_report.php"

from erp_credentials import get_erp_credentials

DEFAULT_CREDENTIALS = get_erp_credentials()

CACHE_FILE = "spare_parts_cache.json"
EXCEL_OUTPUT_FILE = "Fan_Spare_Parts_Stock_Summary.xlsx"

FONT_NAME = "Segoe UI"
COLOR_NAVY = "1B365D"
COLOR_SECONDARY = "2B5B84"
COLOR_LIGHT_BG = "F4F7FB"
COLOR_ZEBRA = "F9FBFC"
COLOR_BORDER = "D1D5DB"

thin_border = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_BORDER),
    bottom=Side(style='thin', color=COLOR_BORDER)
)

header_border = Border(
    left=Side(style='thin', color=COLOR_NAVY),
    right=Side(style='thin', color=COLOR_NAVY),
    top=Side(style='medium', color=COLOR_NAVY),
    bottom=Side(style='medium', color=COLOR_NAVY)
)


class SparePartsCollector:
    def __init__(self, credentials=None, session=None):
        self.creds = credentials or DEFAULT_CREDENTIALS
        if session:
            self.session = session
            self._shared_session = True
        else:
            self.session = requests.Session()
            self.session.headers.update({
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                "Accept-Language": "en-US,en;q=0.9",
            })
            self._shared_session = False
        self.items = []
        self.meta = {}

    def login(self):
        """Authenticates with MEP ERP if not using shared session."""
        if getattr(self, '_shared_session', False):
            return True
        print("[*] Logging into MEP Group ERP...")
        res = self.session.post(LOGIN_URL, data=self.creds, timeout=20)
        if "login" in res.url.lower() and "home.php" not in res.url.lower():
            raise Exception("Authentication failed! Please check credentials.")
        print("[+] Logged in successfully!")
        return True

    def collect_spare_parts(self, f_date_str=None, t_date_str=None):
        """
        Collects Spare parts Stock Summary Report with rate (3224).
        Company: FAN (group_for = 3)
        Default dates: 1st of current month to today.
        """
        now = datetime.now()
        if not f_date_str:
            f_date_str = datetime(now.year, now.month, 1).strftime("%Y-%m-%d")
        if not t_date_str:
            t_date_str = now.strftime("%Y-%m-%d")

        print(f"[*] Querying Spare Parts Stock Summary (3224) from {f_date_str} to {t_date_str}...")

        payload = {
            "report": "3224",
            "group_for": "3",  # FAN
            "f_date": f_date_str,
            "t_date": t_date_str,
            "submit": "Report"
        }

        res = self.session.post(REPORT_URL, data=payload, timeout=40)
        soup = BeautifulSoup(res.text, "html.parser")
        tbl = soup.find("table")
        if not tbl:
            raise Exception("Spare parts report table not found in ERP response.")

        rows = tbl.find_all("tr")
        parsed_items = []

        # Skip banner (row 0) and header (row 1)
        for r in rows[2:]:
            cols = [c.get_text(strip=True) for c in r.find_all(["td", "th"])]
            if len(cols) >= 14 and cols[0].isdigit():
                try:
                    store_qty = float(cols[9].replace(",", "") or 0)
                    sec_qty = float(cols[10].replace(",", "") or 0)
                    tot_qty = float(cols[11].replace(",", "") or 0)
                    rate = float(cols[12].replace(",", "") or 0)
                    stock_amt = float(cols[13].replace(",", "") or 0)
                except ValueError:
                    continue

                parsed_items.append({
                    "sl": cols[0],
                    "company": cols[1],
                    "group": cols[2],
                    "category": cols[3],
                    "subcategory": cols[4],
                    "item_code": cols[5],
                    "fg": cols[6],
                    "item_name": cols[7],
                    "unit": cols[8] or "Pcs",
                    "store_qty": store_qty,
                    "section_qty": sec_qty,
                    "total_qty": tot_qty,
                    "rate": rate,
                    "stock_amt": stock_amt
                })

        self.items = parsed_items
        self.meta = {
            "collected_at": now.strftime("%Y-%m-%d %H:%M:%S"),
            "f_date": f_date_str,
            "t_date": t_date_str,
            "company": "FAN",
            "report_code": "3224",
            "report_name": "Spare parts Stock Summary Report with rate (3224)",
            "total_items": len(parsed_items),
            "in_stock_items": sum(1 for it in parsed_items if it["total_qty"] > 0),
            "total_quantity": sum(it["total_qty"] for it in parsed_items),
            "total_stock_value": sum(it["stock_amt"] for it in parsed_items)
        }

        print(f"[+] Collected {len(parsed_items)} Spare Parts items (Total Value: BDT {self.meta['total_stock_value']:,.2f})")
        return self.items

    def save_cache(self, filepath=CACHE_FILE):
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump({
                "meta": self.meta,
                "items": self.items
            }, f, indent=2, ensure_ascii=False)
        print(f"[+] Saved cache to {filepath}")

    def load_cache(self, filepath=CACHE_FILE):
        if os.path.exists(filepath):
            with open(filepath, "r", encoding="utf-8") as f:
                d = json.load(f)
                self.items = d.get("items", [])
                self.meta = d.get("meta", {})
                return True
        return False

    def export_to_excel(self, output_file=EXCEL_OUTPUT_FILE):
        print(f"[*] Exporting Spare Parts Workbook to {output_file}...")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        # 1. Executive Summary Sheet
        ws_sum = wb.create_sheet(title="Executive Summary")
        ws_sum.views.sheetView[0].showGridLines = True
        self._build_summary_sheet(ws_sum)

        # 2. Master Stock Summary
        ws_main = wb.create_sheet(title="Spare Parts Stock Summary")
        ws_main.views.sheetView[0].showGridLines = True
        self._build_table_sheet(ws_main, self.items, "ALL SPARE PARTS INVENTORY")

        # 3. In-Stock Items Only
        in_stock = [it for it in self.items if it["total_qty"] > 0]
        ws_instock = wb.create_sheet(title="In-Stock Inventory Only")
        ws_instock.views.sheetView[0].showGridLines = True
        self._build_table_sheet(ws_instock, in_stock, "ACTIVE IN-STOCK ITEMS ONLY")

        wb.save(output_file)
        print(f"[+] Successfully generated {output_file}!")
        return output_file

    def _build_summary_sheet(self, ws):
        ws.merge_cells("A1:H2")
        b = ws["A1"]
        b.value = "MEP GROUP  |  FAN Spare Parts Stock Summary with Rate (3224)"
        b.font = Font(name=FONT_NAME, size=15, bold=True, color="FFFFFF")
        b.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        b.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        f_d = self.meta.get("f_date", "")
        t_d = self.meta.get("t_date", "")
        ws["A3"].value = f"Workflow: Warehouse -> Report -> Warehouse Reports -> Report 3224  |  Company: FAN  |  Date Range: {f_d} to {t_d} (Running Date)"
        ws["A3"].font = Font(name=FONT_NAME, size=9.5, italic=True, color="64748B")

        tot_items = self.meta.get("total_items", len(self.items))
        in_stock_items = self.meta.get("in_stock_items", sum(1 for it in self.items if it["total_qty"] > 0))
        tot_qty = self.meta.get("total_quantity", sum(it["total_qty"] for it in self.items))
        tot_val = self.meta.get("total_stock_value", sum(it["stock_amt"] for it in self.items))

        kpis = [
            ("Total Catalog Items", tot_items, "Spare Parts Catalog Registry", "1B365D"),
            ("Active In-Stock Items", in_stock_items, f"{tot_items - in_stock_items} Zero Stock Items", "059669"),
            ("Total In-Stock Qty", f"{tot_qty:,.2f} Pcs", "Cumulative Warehouse Stock", "0284C7"),
            ("Total Stock Valuation", f"BDT {tot_val:,.2f}", "Total Net Inventory Value", "D97706"),
        ]

        col_pairs = [("B", "C"), ("D", "E"), ("F", "G"), ("H", "I")]
        for idx, (lbl, val, sub, col_hex) in enumerate(kpis):
            c1, c2 = col_pairs[idx]
            ws.merge_cells(f"{c1}5:{c2}5")
            ws.merge_cells(f"{c1}6:{c2}6")
            ws.merge_cells(f"{c1}7:{c2}7")

            v_cell = ws[f"{c1}5"]
            v_cell.value = val
            v_cell.font = Font(name=FONT_NAME, size=16, bold=True, color=col_hex)
            v_cell.alignment = Alignment(horizontal="center", vertical="center")

            l_cell = ws[f"{c1}6"]
            l_cell.value = lbl
            l_cell.font = Font(name=FONT_NAME, size=10, bold=True, color="334155")
            l_cell.alignment = Alignment(horizontal="center", vertical="center")

            s_cell = ws[f"{c1}7"]
            s_cell.value = sub
            s_cell.font = Font(name=FONT_NAME, size=8, color="64748B")
            s_cell.alignment = Alignment(horizontal="center", vertical="center")

            for r in range(5, 8):
                for c in [c1, c2]:
                    ws[f"{c}{r}"].fill = PatternFill(start_color=COLOR_LIGHT_BG, end_color=COLOR_LIGHT_BG, fill_type="solid")
                    ws[f"{c}{r}"].border = thin_border

        # SubCategory breakdown table
        ws["B9"].value = "SUBCATEGORY INVENTORY BREAKDOWN"
        ws["B9"].font = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_NAVY)

        subcat_headers = ["SubCategory", "Total Items", "In-Stock Items", "Total Quantity (Pcs)", "Stock Valuation (BDT)", "% Value Share"]
        for c_idx, h in enumerate(subcat_headers, start=2):
            cell = ws.cell(row=10, column=c_idx, value=h)
            cell.font = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            cell.alignment = Alignment(horizontal="left" if c_idx == 2 else "right", vertical="center")
            cell.border = header_border

        subcat_map = {}
        for it in self.items:
            sc = it["subcategory"] or "General / Uncategorized"
            if sc not in subcat_map:
                subcat_map[sc] = {"count": 0, "in_stock": 0, "qty": 0.0, "val": 0.0}
            subcat_map[sc]["count"] += 1
            if it["total_qty"] > 0:
                subcat_map[sc]["in_stock"] += 1
            subcat_map[sc]["qty"] += it["total_qty"]
            subcat_map[sc]["val"] += it["stock_amt"]

        curr_r = 11
        for sc, stats in sorted(subcat_map.items()):
            bg = COLOR_ZEBRA if curr_r % 2 == 1 else "FFFFFF"
            share = (stats["val"] / tot_val * 100) if tot_val > 0 else 0
            row_vals = [sc, stats["count"], stats["in_stock"], stats["qty"], stats["val"], f"{share:.1f}%"]
            for c_idx, v in enumerate(row_vals, start=2):
                c = ws.cell(row=curr_r, column=c_idx, value=v)
                c.font = Font(name=FONT_NAME, size=9.5)
                c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                c.border = thin_border
                c.alignment = Alignment(horizontal="left" if c_idx == 2 else "right", vertical="center")
                if c_idx in [4, 5]:
                    c.number_format = "#,##0.00"
            curr_r += 1

        for col in ["B", "C", "D", "E", "F", "G"]:
            ws.column_dimensions[col].width = 24

    def _build_table_sheet(self, ws, items_list, section_title):
        ws.merge_cells("A1:K1")
        title = ws["A1"]
        title.value = f"MEP GROUP  |  FAN Spare Parts — {section_title}"
        title.font = Font(name=FONT_NAME, size=13, bold=True, color="FFFFFF")
        title.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 28

        headers = [
            ("SL", 8, "center"),
            ("Item Code", 15, "center"),
            ("Part No (FG)", 15, "center"),
            ("Item Name & Description", 38, "left"),
            ("SubCategory", 20, "left"),
            ("Unit", 8, "center"),
            ("Store Qty", 14, "right"),
            ("Section Qty", 14, "right"),
            ("Total Qty", 15, "right"),
            ("Rate (BDT)", 14, "right"),
            ("Stock Value (BDT)", 18, "right")
        ]

        ws.row_dimensions[3].height = 24
        for c_idx, (h_title, _, align) in enumerate(headers, start=1):
            c = ws.cell(row=3, column=c_idx, value=h_title)
            c.font = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            c.alignment = Alignment(horizontal=align, vertical="center")
            c.border = header_border

        curr_r = 4
        for it in items_list:
            bg = COLOR_ZEBRA if curr_r % 2 == 1 else "FFFFFF"
            ws.row_dimensions[curr_r].height = 20
            row_data = [
                it["sl"],
                it["item_code"],
                it["fg"],
                it["item_name"],
                it["subcategory"],
                it["unit"],
                it["store_qty"],
                it["section_qty"],
                it["total_qty"],
                it["rate"],
                it["stock_amt"]
            ]
            for c_idx, v in enumerate(row_data, start=1):
                c = ws.cell(row=curr_r, column=c_idx, value=v)
                c.font = Font(name=FONT_NAME, size=9)
                c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                c.border = thin_border
                c.alignment = Alignment(horizontal=headers[c_idx-1][2], vertical="center")
                if c_idx in [7, 8, 9, 10, 11]:
                    c.number_format = "#,##0.00"
            curr_r += 1

        # Summary Row
        ws.row_dimensions[curr_r].height = 24
        tot_qty = sum(it["total_qty"] for it in items_list)
        tot_val = sum(it["stock_amt"] for it in items_list)

        ws.cell(row=curr_r, column=1, value="Total")
        ws.cell(row=curr_r, column=4, value=f"{len(items_list)} Items")
        ws.cell(row=curr_r, column=9, value=tot_qty)
        ws.cell(row=curr_r, column=11, value=tot_val)

        for c_idx in range(1, 12):
            c = ws.cell(row=curr_r, column=c_idx)
            c.font = Font(name=FONT_NAME, size=10, bold=True, color="000000")
            c.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            c.border = Border(top=Side(style='thin'), bottom=Side(style='double'))
            if c_idx in [9, 11]:
                c.number_format = "#,##0.00"

        for c_idx, (_, width, _) in enumerate(headers, start=1):
            col_letter = get_column_letter(c_idx)
            ws.column_dimensions[col_letter].width = width


if __name__ == "__main__":
    collector = SparePartsCollector()
    collector.login()
    collector.collect_spare_parts()
    collector.save_cache()
    collector.export_to_excel()
    print("Spare parts collection completed successfully.")
