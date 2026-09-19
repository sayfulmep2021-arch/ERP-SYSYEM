"""
MEP Group ERP - Inter Sales Requisition Collector
Module: Warehouse Module -> Requisition Status -> Company To: FAN -> View Detail
Special Date Requirement: User-specified From Date and To Date (From Date <= Requisition Date <= To Date)
Author: Antigravity
"""

import os
import re
import json
import time
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from erp_credentials import get_erp_credentials

BASE_URL = "https://www.mepgrouperp.com/1027"
LOGIN_URL = f"{BASE_URL}/login/pages/main/index.php"
MR_STATUS_URL = f"{BASE_URL}/warehouse_mod/pages/inter_sales/mr_status.php"
PRINT_VIEW_URL = f"{BASE_URL}/warehouse_mod/pages/inter_sales/intersales_print_view.php?req_no="

CACHE_FILE = "intersales_requisition_cache.json"
EXCEL_OUTPUT_FILE = "Inter_Sales_Requisition_Report.xlsx"

# Visual Styles for Premium Excel
FONT_NAME = "Segoe UI"
COLOR_NAVY = "1B365D"       # Corporate Dark Navy
COLOR_SECONDARY = "2B5B84"  # Medium Slate Navy
COLOR_ACCENT = "0284C7"     # Electric Ocean Blue
COLOR_LIGHT_BG = "F4F7FB"   # Soft Blue-Gray
COLOR_ZEBRA = "EAF5EA"      # Subtle Paste Row
COLOR_BORDER = "D1D5DB"     # Soft Gray
COLOR_ACTIVE_BG = "E6F4EA"  # Soft Green
COLOR_ACTIVE_FG = "137333"  # Dark Green
COLOR_PENDING_BG = "FEF3C7" # Soft Amber
COLOR_PENDING_FG = "B45309" # Dark Amber

thin_border = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_BORDER),
    bottom=Side(style='thin', color=COLOR_BORDER)
)

header_border = Border(
    left=Side(style='thin', color=COLOR_NAVY),
    right=Side(style='thin', color=COLOR_NAVY),
    top=Side(style='medium', color=COLOR_NAVY),
    bottom=Side(style='medium', color=COLOR_NAVY)
)


def parse_date_safely(date_str):
    """Parses various date string formats into a datetime.date object."""
    if not date_str:
        return None
    cleaned = date_str.strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%b-%Y"):
        try:
            return datetime.strptime(cleaned, fmt).date()
        except ValueError:
            pass
    return None


class InterSalesRequisitionCollector:
    def __init__(self, credentials=None, session=None):
        self.creds = credentials or get_erp_credentials()
        if session:
            self.session = session
            self._shared_session = True
        else:
            self.session = requests.Session()
            self.session.headers.update({
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                "Accept-Language": "en-US,en;q=0.9",
            })
            self._shared_session = False
        self.requisitions = []
        self.meta = {}

    def login(self):
        """Authenticates with MEP ERP if not already using a shared authenticated session."""
        if getattr(self, '_shared_session', False):
            return True
        print("[*] Logging into MEP Group ERP for Inter Sales Requisition...")
        res = self.session.post(LOGIN_URL, data=self.creds, allow_redirects=True, timeout=20)
        if "login" in res.url.lower() and "home.php" not in res.url.lower():
            raise Exception("Authentication failed! Please verify credentials.")
        print("[+] Logged in successfully!")
        return True

    def collect(self, f_date_str=None, t_date_str=None, company_to_id="3"):
        """
        Collects Inter Sales Requisitions for Company To: FAN (id: 3)
        Within the user-specified date range (From Date <= Requisition Date <= To Date).
        """
        now = datetime.now()
        if not f_date_str:
            f_date_str = datetime(now.year, now.month, 1).strftime("%Y-%m-%d")
        if not t_date_str:
            t_date_str = now.strftime("%Y-%m-%d")

        f_date_obj = parse_date_safely(f_date_str)
        t_date_obj = parse_date_safely(t_date_str)

        # Standardize format for ERP POST (YYYY-MM-DD)
        erp_fdate = f_date_obj.strftime("%Y-%m-%d") if f_date_obj else f_date_str
        erp_tdate = t_date_obj.strftime("%Y-%m-%d") if t_date_obj else t_date_str

        print(f"[*] Querying Requisition Status from {erp_fdate} to {erp_tdate} (Company To: FAN)...")
        payload = {
            "fdate": erp_fdate,
            "tdate": erp_tdate,
            "group_for": company_to_id,
            "submitit": "VIEW DETAIL"
        }

        res = self.session.post(MR_STATUS_URL, data=payload, timeout=25)
        if res.status_code != 200:
            raise Exception(f"Failed to load Requisition Status page (HTTP {res.status_code})")

        soup = BeautifulSoup(res.text, "html.parser")
        tbl = soup.find("table", id="grp")

        parsed_reqs = []
        if tbl:
            rows = tbl.find_all("tr")
            # First row is header: ['Req No', 'Req Date', 'Company For', 'Company To', 'Need By', 'Entry By', 'Status']
            for tr in rows[1:]:
                cols = [c.get_text(strip=True) for c in tr.find_all(["th", "td"])]
                if len(cols) >= 7 and cols[0] and cols[0] != "Req No":
                    req_no = cols[0]
                    req_date = cols[1]
                    company_for = cols[2]
                    company_to = cols[3]
                    need_by = cols[4]
                    entry_by = cols[5]
                    status = cols[6]

                    # Enforce strict date range restriction: From Date <= Requisition Date <= To Date
                    req_date_obj = parse_date_safely(req_date)
                    if req_date_obj and f_date_obj and t_date_obj:
                        if not (f_date_obj <= req_date_obj <= t_date_obj):
                            print(f"[-] Skipping out-of-range Requisition {req_no} ({req_date})")
                            continue

                    parsed_reqs.append({
                        "req_no": req_no,
                        "req_date": req_date,
                        "company_for": company_for,
                        "company_to": company_to,
                        "need_by": need_by,
                        "entry_by": entry_by,
                        "status": status,
                        "items": [],
                        "total_req_qty": 0.0,
                        "total_issue_qty": 0.0,
                        "total_pending_qty": 0.0,
                        "materials_count": 0
                    })

        print(f"[+] Found {len(parsed_reqs)} Requisitions within date range. Fetching detailed items...")

        # Fetch print view details concurrently
        def fetch_detail(req_dict):
            req_no = req_dict["req_no"]
            try:
                pv_url = f"{PRINT_VIEW_URL}{req_no}"
                pv_res = self.session.get(pv_url, timeout=15)
                if pv_res.status_code == 200:
                    pv_soup = BeautifulSoup(pv_res.text, "html.parser")

                    # Extract header info if present
                    text = pv_soup.get_text()
                    from_wh_match = re.search(r"From Warehouse\s*:\s*([^<\n]+)", text)
                    to_wh_match = re.search(r"To Warehouse\s*:\s*([^<\n]+)", text)
                    if from_wh_match:
                        req_dict["from_warehouse"] = from_wh_match.group(1).strip()
                    if to_wh_match:
                        req_dict["to_warehouse"] = to_wh_match.group(1).strip()

                    for pvtbl in pv_soup.find_all("table"):
                        prows = pvtbl.find_all("tr")
                        if len(prows) > 1:
                            hdr_cols = [c.get_text(strip=True).lower() for c in prows[0].find_all(["th", "td"])]
                            if any("code" in h for h in hdr_cols) and any("desc" in h or "goods" in h for h in hdr_cols):
                                for pr in prows[1:]:
                                    pcols = [c.get_text(strip=True) for c in pr.find_all(["th", "td"])]
                                    if len(pcols) >= 6:
                                        sl = pcols[0]
                                        item_code = pcols[1]
                                        desc = pcols[2]
                                        unit = pcols[3] if len(pcols) > 3 else "Pcs"
                                        try:
                                            req_qty = float(pcols[4].replace(",", "")) if pcols[4] else 0.0
                                        except ValueError:
                                            req_qty = 0.0
                                        try:
                                            issue_qty = float(pcols[5].replace(",", "")) if pcols[5] else 0.0
                                        except ValueError:
                                            issue_qty = 0.0
                                        pending_qty = 0.0
                                        if len(pcols) > 6 and pcols[6]:
                                            try:
                                                pending_qty = float(pcols[6].replace(",", ""))
                                            except ValueError:
                                                pending_qty = req_qty - issue_qty
                                        else:
                                            pending_qty = max(0.0, req_qty - issue_qty)

                                        req_dict["items"].append({
                                            "sl": sl,
                                            "item_code": item_code,
                                            "item_name": desc,
                                            "unit": unit,
                                            "req_qty": req_qty,
                                            "issue_qty": issue_qty,
                                            "pending_qty": pending_qty
                                        })
            except Exception as e:
                print(f"[!] Error fetching items for req {req_no}: {e}")

            req_dict["total_req_qty"] = sum(it["req_qty"] for it in req_dict["items"])
            req_dict["total_issue_qty"] = sum(it["issue_qty"] for it in req_dict["items"])
            req_dict["total_pending_qty"] = sum(it["pending_qty"] for it in req_dict["items"])
            req_dict["materials_count"] = len(req_dict["items"])
            return req_dict

        with ThreadPoolExecutor(max_workers=5) as executor:
            self.requisitions = list(executor.map(fetch_detail, parsed_reqs))

        # Sort by requisition date descending, then req_no descending
        self.requisitions.sort(key=lambda r: (r["req_date"], r["req_no"]), reverse=True)

        total_req_qty = sum(r["total_req_qty"] for r in self.requisitions)
        total_issue_qty = sum(r["total_issue_qty"] for r in self.requisitions)
        total_pending_qty = sum(r["total_pending_qty"] for r in self.requisitions)
        total_items = sum(r["materials_count"] for r in self.requisitions)
        pending_reqs = sum(1 for r in self.requisitions if "PEND" in (r.get("status") or "").upper())

        self.meta = {
            "title": "Inter Sales Requisition",
            "full_title": "Inter Sales Requisition (Warehouse -> Requisition Status -> FAN)",
            "f_date": erp_fdate,
            "t_date": erp_tdate,
            "company_to": "FAN",
            "total_requisitions": len(self.requisitions),
            "pending_requisitions": pending_reqs,
            "total_items": total_items,
            "total_req_qty": round(total_req_qty, 2),
            "total_issue_qty": round(total_issue_qty, 2),
            "total_pending_qty": round(total_pending_qty, 2),
            "collected_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }

        print(f"[+] Successfully collected {len(self.requisitions)} requisitions ({total_items} items total)")
        return self.requisitions

    def save_cache(self, filepath=None):
        target = filepath or CACHE_FILE
        data = {
            "meta": self.meta,
            "requisitions": self.requisitions
        }
        with open(target, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2)
        print(f"[+] Saved cache to {target}")

    def load_cache(self, filepath=None):
        target = filepath or CACHE_FILE
        if os.path.exists(target):
            try:
                with open(target, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    self.meta = data.get("meta", {})
                    self.requisitions = data.get("requisitions", [])
                return True
            except Exception as e:
                print(f"[!] Failed to load cache from {target}: {e}")
        return False

    def export_to_excel(self, filepath=None):
        """Generates executive openpyxl workbook with alternating paste colors."""
        target = filepath or EXCEL_OUTPUT_FILE
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Requisition Summary"
        ws.views.sheetView[0].showGridLines = True

        # Sheet 1: Summary List
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 22
        ws.column_dimensions["D"].width = 24
        ws.column_dimensions["E"].width = 14
        ws.column_dimensions["F"].width = 20
        ws.column_dimensions["G"].width = 14
        ws.column_dimensions["H"].width = 14
        ws.column_dimensions["I"].width = 14
        ws.column_dimensions["J"].width = 14
        ws.column_dimensions["K"].width = 12

        # Title Block
        ws.merge_cells("A1:K1")
        top_cell = ws["A1"]
        top_cell.value = "MEP GROUP ERP — INTER SALES REQUISITION STATUS REPORT"
        top_cell.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
        top_cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        top_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        # Subtitle
        ws.merge_cells("A2:K2")
        sub_cell = ws["A2"]
        f_date = self.meta.get("f_date", "")
        t_date = self.meta.get("t_date", "")
        total_reqs = self.meta.get("total_requisitions", len(self.requisitions))
        sub_cell.value = f"Destination: FAN  |  Date Range: {f_date} to {t_date}  |  Total Requisitions: {total_reqs}  |  Generated: {self.meta.get('collected_at', '')}"
        sub_cell.font = Font(name=FONT_NAME, size=9.5, italic=True, color="1E293B")
        sub_cell.fill = PatternFill(start_color=COLOR_LIGHT_BG, end_color=COLOR_LIGHT_BG, fill_type="solid")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        headers = [
            "Req No", "Req Date", "Company For", "Company To",
            "Need By", "Entry By", "Status", "Total Req Qty",
            "Total Issue Qty", "Pending Qty", "Items Count"
        ]

        ws.row_dimensions[4].height = 26
        for col_num, h_title in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num, value=h_title)
            cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=COLOR_SECONDARY, end_color=COLOR_SECONDARY, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = header_border

        row_idx = 5
        for r in self.requisitions:
            bg_color = COLOR_ZEBRA if (row_idx % 2 == 0) else "FFFFFF"
            fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

            ws.cell(row=row_idx, column=1, value=r.get("req_no", "")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=r.get("req_date", "")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=3, value=r.get("company_for", "")).alignment = Alignment(horizontal="left")
            ws.cell(row=row_idx, column=4, value=r.get("company_to", "")).alignment = Alignment(horizontal="left")
            ws.cell(row=row_idx, column=5, value=r.get("need_by", "")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=6, value=r.get("entry_by", "")).alignment = Alignment(horizontal="left")

            st_cell = ws.cell(row=row_idx, column=7, value=r.get("status", ""))
            st_cell.alignment = Alignment(horizontal="center")
            if "PEND" in (r.get("status") or "").upper():
                st_cell.fill = PatternFill(start_color=COLOR_PENDING_BG, end_color=COLOR_PENDING_BG, fill_type="solid")
                st_cell.font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_PENDING_FG)
            else:
                st_cell.fill = PatternFill(start_color=COLOR_ACTIVE_BG, end_color=COLOR_ACTIVE_BG, fill_type="solid")
                st_cell.font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_ACTIVE_FG)

            c8 = ws.cell(row=row_idx, column=8, value=r.get("total_req_qty", 0))
            c8.number_format = "#,##0.00"
            c8.alignment = Alignment(horizontal="right")

            c9 = ws.cell(row=row_idx, column=9, value=r.get("total_issue_qty", 0))
            c9.number_format = "#,##0.00"
            c9.alignment = Alignment(horizontal="right")

            c10 = ws.cell(row=row_idx, column=10, value=r.get("total_pending_qty", 0))
            c10.number_format = "#,##0.00"
            c10.alignment = Alignment(horizontal="right")

            c11 = ws.cell(row=row_idx, column=11, value=r.get("materials_count", 0))
            c11.number_format = "#,##0"
            c11.alignment = Alignment(horizontal="center")

            for c in range(1, 12):
                cell = ws.cell(row=row_idx, column=c)
                if c != 7:
                    cell.fill = fill
                    cell.font = Font(name=FONT_NAME, size=9.5)
                cell.border = thin_border

            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        # Summary Row
        ws.row_dimensions[row_idx].height = 24
        ws.merge_cells(f"A{row_idx}:G{row_idx}")
        sum_label = ws.cell(row=row_idx, column=1, value="TOTAL SUMMARY")
        sum_label.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        sum_label.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        sum_label.alignment = Alignment(horizontal="center", vertical="center")

        t_req = self.meta.get("total_req_qty", 0)
        t_iss = self.meta.get("total_issue_qty", 0)
        t_pen = self.meta.get("total_pending_qty", 0)
        t_itm = self.meta.get("total_items", 0)

        for c_idx, val in [(8, t_req), (9, t_iss), (10, t_pen), (11, t_itm)]:
            c = ws.cell(row=row_idx, column=c_idx, value=val)
            c.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            c.border = header_border
            c.number_format = "#,##0.00" if c_idx < 11 else "#,##0"
            c.alignment = Alignment(horizontal="right" if c_idx < 11 else "center", vertical="center")

        # Sheet 2: Itemized Materials Detail
        ws_items = wb.create_sheet(title="Itemized Materials")
        ws_items.views.sheetView[0].showGridLines = True

        ws_items.column_dimensions["A"].width = 14
        ws_items.column_dimensions["B"].width = 14
        ws_items.column_dimensions["C"].width = 8
        ws_items.column_dimensions["D"].width = 16
        ws_items.column_dimensions["E"].width = 38
        ws_items.column_dimensions["F"].width = 10
        ws_items.column_dimensions["G"].width = 14
        ws_items.column_dimensions["H"].width = 14
        ws_items.column_dimensions["I"].width = 14

        item_headers = ["Req No", "Req Date", "SL", "Item Code", "Item Description", "Unit", "Req Qty", "Issue Qty", "Pending Qty"]
        ws_items.row_dimensions[1].height = 26
        for col_num, h_title in enumerate(item_headers, 1):
            cell = ws_items.cell(row=1, column=col_num, value=h_title)
            cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = header_border

        i_row = 2
        for r in self.requisitions:
            for item in r.get("items", []):
                bg = COLOR_ZEBRA if (i_row % 2 == 0) else "FFFFFF"
                fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")

                ws_items.cell(row=i_row, column=1, value=r.get("req_no", "")).alignment = Alignment(horizontal="center")
                ws_items.cell(row=i_row, column=2, value=r.get("req_date", "")).alignment = Alignment(horizontal="center")
                ws_items.cell(row=i_row, column=3, value=item.get("sl", "")).alignment = Alignment(horizontal="center")
                ws_items.cell(row=i_row, column=4, value=item.get("item_code", "")).alignment = Alignment(horizontal="center")
                ws_items.cell(row=i_row, column=5, value=item.get("item_name", "")).alignment = Alignment(horizontal="left")
                ws_items.cell(row=i_row, column=6, value=item.get("unit", "")).alignment = Alignment(horizontal="center")

                c7 = ws_items.cell(row=i_row, column=7, value=item.get("req_qty", 0))
                c7.number_format = "#,##0.00"
                c7.alignment = Alignment(horizontal="right")

                c8 = ws_items.cell(row=i_row, column=8, value=item.get("issue_qty", 0))
                c8.number_format = "#,##0.00"
                c8.alignment = Alignment(horizontal="right")

                c9 = ws_items.cell(row=i_row, column=9, value=item.get("pending_qty", 0))
                c9.number_format = "#,##0.00"
                c9.alignment = Alignment(horizontal="right")

                for c in range(1, 10):
                    cell = ws_items.cell(row=i_row, column=c)
                    cell.fill = fill
                    cell.font = Font(name=FONT_NAME, size=9.5)
                    cell.border = thin_border

                ws_items.row_dimensions[i_row].height = 19
                i_row += 1

        wb.save(target)
        print(f"[+] Saved Excel to {target}")
        return target


if __name__ == "__main__":
    collector = InterSalesRequisitionCollector()
    collector.login()
    collector.collect("2026-09-01", "2026-09-17")
    collector.save_cache()
    collector.export_to_excel()
