"""
MEP Group ERP - Monthly & Yearly Attendance Report Collector
Path: HRM Module ⟶ New-HRM Report ⟶ Reports ⟶ Monthly & Yearly Attendence Report ⟶ SHOW
Filters: Company: MEP FAN LIMITED. (PBI_ORG=3), Department: Production (DEPT_ID=32)
Report ID: 30082026
Date Range: From Date & To Date (Persistent Configuration)
Author: Antigravity
"""

import os
import re
import json
import time
from datetime import datetime

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from erp_credentials import get_erp_credentials, get_shared_erp_session

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CACHE_FILE = os.path.join(BASE_DIR, "monthly_yearly_attendance_cache.json")
EXCEL_OUTPUT_FILE = os.path.join(BASE_DIR, "Monthly_Yearly_Attendance_Report.xlsx")

BASE_URL = "https://mvc.mepgrouperp.com"
LOGIN_URL = f"{BASE_URL}/app/views/auth/masters/"
REPORT_POST_URL = f"{BASE_URL}/app/views/hrm_mod/report/master_report_att_management.php"

# Visual Styles for Premium Excel
FONT_NAME = "Segoe UI"
COLOR_NAVY = "1B365D"       # Corporate Dark Navy
COLOR_SECONDARY = "2B5B84"  # Medium Slate Navy
COLOR_TEAL = "0D9488"       # Teal
COLOR_LIGHT_BG = "F4F7FB"   # Soft Blue-Gray
COLOR_ZEBRA = "EAF5EA"      # Subtle Paste Row
COLOR_BORDER = "D1D5DB"     # Soft Gray

thin_border = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_BORDER),
    bottom=Side(style='thin', color=COLOR_BORDER)
)

header_border = Border(
    left=Side(style='thin', color=COLOR_NAVY),
    right=Side(style='thin', color=COLOR_NAVY),
    top=Side(style='medium', color=COLOR_NAVY),
    bottom=Side(style='medium', color=COLOR_NAVY)
)


def format_date_for_erp(date_str):
    """Converts DD-MM-YYYY or any standard date to YYYY-MM-DD for ERP POST."""
    if not date_str:
        return "2026-09-01"
    cleaned = str(date_str).strip()
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            dt = datetime.strptime(cleaned, fmt)
            return dt.strftime("%Y-%m-%d")
        except ValueError:
            pass
    return cleaned


class MonthlyYearlyAttendanceCollector:
    def __init__(self, credentials=None, session=None):
        self.creds = credentials or get_erp_credentials()
        self.session = session or requests.Session()
        self.headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
        }
        self.session.headers.update(self.headers)
        self.is_authenticated = False
        self.records = []
        self.column_headers = []
        self.meta = {}

    def login(self):
        """Authenticates against ERP session."""
        try:
            login_data = {
                "db": self.creds.get("db", "erpcombd"),
                "cid": self.creds.get("cid", "MEP"),
                "uid": self.creds.get("uid", "10676"),
                "ibssignin": "",
                "pass": self.creds.get("pass", "96420132"),
                "submit": "Sign In"
            }
            res = self.session.post(LOGIN_URL, data=login_data, timeout=25, allow_redirects=True)
            if "login" in res.url.lower() and "home.php" not in res.url.lower():
                raise Exception("MEP ERP Login failed! Check credentials.")
            self.is_authenticated = True
            return True
        except Exception as e:
            print(f"[!] Login exception: {e}")
            raise

    def collect(self, f_date_str="01-09-2026", t_date_str="21-09-2026", company_id="3", dept_id="32"):
        """
        Submits request for Report 30082026 (Monthly & Yearly Attendence Report)
        Strictly filtered by:
          Company: MEP FAN LIMITED. (PBI_ORG=3)
          Department: Production (DEPT_ID=32)
        with specified From Date and To Date.
        """
        if not self.is_authenticated:
            self.login()

        f_erp = format_date_for_erp(f_date_str)
        t_erp = format_date_for_erp(t_date_str)
        print(f"[*] Monthly & Yearly Attendance Collector: Requesting range {f_erp} to {t_erp} (Company: MEP FAN LIMITED., Dept: Production)...")

        payload = {
            "report": "30082026",
            "PBI_ORG": company_id or "3",
            "DEPT_ID": dept_id or "32",
            "PBI_JOB_STATUS": "In Service",
            "fdate": f_erp,
            "tdate": t_erp,
            "submit": "SHOW"
        }

        res = self.session.post(REPORT_POST_URL, data=payload, timeout=75)
        if res.status_code != 200:
            raise Exception(f"ERP HTTP Error: Status {res.status_code}")

        soup = BeautifulSoup(res.text, "html.parser")
        tables = soup.find_all("table")
        if len(tables) < 2:
            print(f"[!] Warning: Expected at least 2 tables, found {len(tables)}.")
            if len(tables) == 1:
                data_table = tables[0]
            else:
                raise Exception("No tables found in ERP response.")
        else:
            data_table = tables[1]

        rows = data_table.find_all("tr")
        if not rows:
            raise Exception("Monthly & Yearly Attendance table contains no rows.")

        raw_headers = [c.text.strip().replace("\n", " ") for c in rows[0].find_all(["th", "td"])]
        self.column_headers = raw_headers

        def is_section_allowed(sec_name):
            if not sec_name:
                return False
            norm = re.sub(r'[^a-z0-9]', '', str(sec_name).lower())
            return norm in {
                "assembleline",
                "dimmerblade",
                "dimmarblade",
                "armaturewinding"
            }

        parsed_records = []
        total_present_days = 0
        total_absent_days = 0
        total_ot_hours = 0.0

        current_emp = {
            "sl": "1", "emp_id": "", "emp_name": "", "designation": "",
            "department": "Production", "section": "", "sub_section": "",
            "grade": "", "job_location": "", "company": "MEP FAN LIMITED.", "join_date": ""
        }
        skip_current_emp = False
        emp_counter = 0

        for r_idx in range(1, len(rows)):
            tds = rows[r_idx].find_all(["td", "th"])
            if not tds or len(tds) < 5:
                continue

            vals = [td.text.strip().replace("\xa0", " ") for td in tds]

            if len(vals) >= 41:
                sec = vals[5] if len(vals) > 5 else ""
                if not is_section_allowed(sec):
                    skip_current_emp = True
                    continue

                skip_current_emp = False
                emp_counter += 1

                # Primary row with employee identity
                current_emp = {
                    "sl": str(emp_counter),
                    "emp_id": vals[1] if len(vals) > 1 else "",
                    "emp_name": vals[2] if len(vals) > 2 else "",
                    "designation": vals[3] if len(vals) > 3 else "",
                    "department": vals[4] if len(vals) > 4 else "Production",
                    "section": sec,
                    "sub_section": vals[6] if len(vals) > 6 else "",
                    "grade": vals[7] if len(vals) > 7 else "",
                    "job_location": vals[8] if len(vals) > 8 else "",
                    "company": vals[9] if len(vals) > 9 else "MEP FAN LIMITED.",
                    "join_date": vals[10] if len(vals) > 10 else "",
                }
                month_vals = vals[11:]
                is_first = True
            else:
                if skip_current_emp:
                    continue
                # Sub-row for same employee (due to HTML rowspan on first 11 columns)
                month_vals = vals
                is_first = False

            # Ensure month_vals has 30 elements
            while len(month_vals) < 30:
                month_vals.append("0")

            rec = {
                "sl": current_emp["sl"],
                "emp_id": current_emp["emp_id"],
                "emp_name": current_emp["emp_name"],
                "designation": current_emp["designation"],
                "department": current_emp["department"],
                "section": current_emp["section"],
                "sub_section": current_emp["sub_section"],
                "grade": current_emp["grade"],
                "job_location": current_emp["job_location"],
                "company": current_emp["company"],
                "join_date": current_emp["join_date"],
                "month_year": month_vals[0] if len(month_vals) > 0 else "",
                "assigned_working_hour": month_vals[1] if len(month_vals) > 1 else "0",
                "total_days_in_month": month_vals[2] if len(month_vals) > 2 else "0",
                "festival_holiday": month_vals[3] if len(month_vals) > 3 else "0",
                "total_working_days": month_vals[4] if len(month_vals) > 4 else "0",
                "standard_working_hour": month_vals[5] if len(month_vals) > 5 else "0",
                "present_days": month_vals[6] if len(month_vals) > 6 else "0",
                "physical_working_hour": month_vals[7] if len(month_vals) > 7 else "0",
                "absent_days": month_vals[8] if len(month_vals) > 8 else "0",
                "lwp": month_vals[9] if len(month_vals) > 9 else "0",
                "cl": month_vals[10] if len(month_vals) > 10 else "0",
                "ml": month_vals[11] if len(month_vals) > 11 else "0",
                "el": month_vals[12] if len(month_vals) > 12 else "0",
                "npl": month_vals[13] if len(month_vals) > 13 else "0",
                "paid_leave": month_vals[14] if len(month_vals) > 14 else "0",
                "ot": month_vals[15] if len(month_vals) > 15 else "0",
                "extra_ot": month_vals[16] if len(month_vals) > 16 else "0",
                "total_ot": month_vals[17] if len(month_vals) > 17 else "0",
                "late_days": month_vals[18] if len(month_vals) > 18 else "0",
                "late_min": month_vals[19] if len(month_vals) > 19 else "0",
                "early_days": month_vals[20] if len(month_vals) > 20 else "0",
                "early_min": month_vals[21] if len(month_vals) > 21 else "0",
                "compensatory_leave": month_vals[22] if len(month_vals) > 22 else "0",
                "sp_iom_pay": month_vals[23] if len(month_vals) > 23 else "0",
                "sp_iom_leave": month_vals[24] if len(month_vals) > 24 else "0",
                "reg_iom": month_vals[25] if len(month_vals) > 25 else "0",
                "od_iom": month_vals[26] if len(month_vals) > 26 else "0",
                "night_duty_days": month_vals[27] if len(month_vals) > 27 else "0",
                "night_duty_hours": month_vals[28] if len(month_vals) > 28 else "0",
                "m_grade_extra_duty": month_vals[29] if len(month_vals) > 29 else "0",
                "is_first_month_row": is_first,
                "all_cells": [
                    current_emp["sl"], current_emp["emp_id"], current_emp["emp_name"],
                    current_emp["designation"], current_emp["department"], current_emp["section"],
                    current_emp["sub_section"], current_emp["grade"], current_emp["job_location"],
                    current_emp["company"], current_emp["join_date"]
                ] + month_vals
            }

            try:
                total_present_days += float(rec["present_days"] or 0)
            except ValueError:
                pass
            try:
                total_absent_days += float(rec["absent_days"] or 0)
            except ValueError:
                pass
            try:
                total_ot_hours += float(rec["total_ot"] or 0)
            except ValueError:
                pass

            parsed_records.append(rec)

        self.records = parsed_records
        self.meta = {
            "report_name": "Monthly & Yearly Attendence Report",
            "report_id": "30082026",
            "source_path": "HRM Module ⟶ New-HRM Report ⟶ Reports ⟶ Monthly & Yearly Attendence Report ⟶ SHOW",
            "filter_company": "MEP FAN LIMITED.",
            "filter_department": "Production",
            "allowed_sections": ["Assemble Line", "Dimmer & Blade", "Armature & winding"],
            "company_id": company_id,
            "dept_id": dept_id,
            "from_date": f_date_str,
            "to_date": t_date_str,
            "fdate_erp": f_erp,
            "tdate_erp": t_erp,
            "total_employees": emp_counter,
            "total_records": len(parsed_records),
            "total_present_days": round(total_present_days, 1),
            "total_absent_days": round(total_absent_days, 1),
            "total_ot_hours": round(total_ot_hours, 1),
            "columns_count": len(self.column_headers),
            "collected_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }

        print(f"[+] Successfully extracted {len(parsed_records)} records across {emp_counter} employees for allowed sections.")
        return self.records

    def save_cache(self, path=None):
        """Saves collected data to JSON cache with Sync Failure Protection."""
        target_path = path or CACHE_FILE
        if not self.records and os.path.exists(target_path):
            print("[!] Empty records returned. Preserving existing valid cache!")
            return False

        payload = {
            "meta": self.meta,
            "headers": self.column_headers,
            "items": self.records,
            "summary": {
                "total_records": len(self.records),
                "total_present_days": self.meta.get("total_present_days", 0),
                "total_absent_days": self.meta.get("total_absent_days", 0),
                "total_ot_hours": self.meta.get("total_ot_hours", 0),
                "filter_company": "MEP FAN LIMITED.",
                "filter_department": "Production",
                "last_sync": datetime.now().strftime("%Y-%m-%d %I:%M %p")
            }
        }
        with open(target_path, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2, ensure_ascii=False)
        print(f"[+] Saved Monthly & Yearly Attendance cache to: {target_path}")
        return True

    def export_to_excel(self, output_path=None):
        """Generates an executive-styled multi-column Excel workbook matching ERP exact structure."""
        out_file = output_path or EXCEL_OUTPUT_FILE
        if not self.records:
            return False

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Monthly & Yearly Attendance"
        ws.views.sheetView[0].showGridLines = True

        # 1. Header Banner
        ws.merge_cells("A1:K1")
        title_cell = ws["A1"]
        title_cell.value = "MEP GROUP — MONTHLY & YEARLY ATTENDANCE REPORT"
        title_cell.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        # Subtitle
        ws.merge_cells("A2:K2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Company: MEP FAN LIMITED. | Department: Production | Period: {self.meta.get('from_date')} to {self.meta.get('to_date')} | Total Records: {len(self.records)}"
        sub_cell.font = Font(name=FONT_NAME, size=10, italic=True, color="FFFFFF")
        sub_cell.fill = PatternFill(start_color=COLOR_SECONDARY, end_color=COLOR_SECONDARY, fill_type="solid")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        # 2. Table Headers
        headers = self.column_headers if self.column_headers else [
            "SL", "Emp ID", "Emp Name", "Designation", "Department", "Section", "Sub-Section", "Grade", "Job Location", "Company",
            "Join Date", "Month-Year", "Assigned working hour", "Total Days in Month", "Festival/Weekly Holiday", "Total Working Days",
            "Standard Working Hour", "Present Days", "Physical Working Hour", "Absent Days", "LWP", "CL", "ML", "EL", "NPL", "Paid Leave",
            "OT", "Extra OT", "Total OT", "Late Days", "Late Min", "Early Days", "Early Min", "Compensatory Leave", "SP IOM Pay",
            "SP IOM Leave", "Reg. IOM", "OD IOM", "Assaigned Night Duty Days", "Assaigned Night Duty Hours", "M-Grade Extra Duty Days"
        ]

        ws.append([]) # Row 3 blank
        ws.append(headers) # Row 4
        header_row_idx = 4
        ws.row_dimensions[header_row_idx].height = 28

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.font = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = header_border

        # 3. Data Rows
        current_row = 5
        for it in self.records:
            row_data = it.get("all_cells", [])
            if not row_data:
                row_data = [
                    it.get("sl", ""), it.get("emp_id", ""), it.get("emp_name", ""), it.get("designation", ""),
                    it.get("department", ""), it.get("section", ""), it.get("sub_section", ""), it.get("grade", ""),
                    it.get("job_location", ""), it.get("company", ""), it.get("join_date", ""), it.get("month_year", ""),
                    it.get("assigned_working_hour", ""), it.get("total_days_in_month", ""), it.get("festival_holiday", ""),
                    it.get("total_working_days", ""), it.get("standard_working_hour", ""), it.get("present_days", ""),
                    it.get("physical_working_hour", ""), it.get("absent_days", ""), it.get("lwp", ""), it.get("cl", ""),
                    it.get("ml", ""), it.get("el", ""), it.get("npl", ""), it.get("paid_leave", ""), it.get("ot", ""),
                    it.get("extra_ot", ""), it.get("total_ot", ""), it.get("late_days", ""), it.get("late_min", ""),
                    it.get("early_days", ""), it.get("early_min", ""), it.get("compensatory_leave", ""), it.get("sp_iom_pay", ""),
                    it.get("sp_iom_leave", ""), it.get("reg_iom", ""), it.get("od_iom", ""), it.get("night_duty_days", ""),
                    it.get("night_duty_hours", ""), it.get("m_grade_extra_duty", "")
                ]

            ws.append(row_data)
            ws.row_dimensions[current_row].height = 20
            fill_color = COLOR_ZEBRA if (current_row % 2 == 0) else "FFFFFF"

            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = Font(name=FONT_NAME, size=9)
                cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", horizontal="left" if col_idx in [3, 4, 5, 6] else "center")

            current_row += 1

        # Auto column dimensions
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 28), 9)

        wb.save(out_file)
        print(f"[+] Saved Monthly & Yearly Attendance Excel to: {out_file}")
        return True


if __name__ == "__main__":
    collector = MonthlyYearlyAttendanceCollector()
    collector.login()
    collector.collect("01-09-2026", "21-09-2026")
    collector.save_cache()
    collector.export_to_excel()
