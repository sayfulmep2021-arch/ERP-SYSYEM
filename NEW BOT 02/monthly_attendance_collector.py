"""
MEP Group ERP - Monthly Attendance Sheet Collector
Path: Login ⟶ HRM Module ⟶ New-HRM Report ⟶ Reports ⟶ Monthly Attendence Sheet ⟶ SHOW
Report ID: 20220522
Date Range: From Date & To Date (Persistent Configuration)
Author: Antigravity
"""

import os
import re
import json
import time
from datetime import datetime

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from erp_credentials import get_erp_credentials, get_shared_erp_session

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CACHE_FILE = os.path.join(BASE_DIR, "monthly_attendance_cache.json")
EXCEL_OUTPUT_FILE = os.path.join(BASE_DIR, "Monthly_Attendance_Sheet.xlsx")

BASE_URL = "https://mvc.mepgrouperp.com"
LOGIN_URL = f"{BASE_URL}/app/views/auth/masters/"
REPORT_PAGE_URL = f"{BASE_URL}/app/views/hrm_mod/report/daily_reports.php"
REPORT_POST_URL = f"{BASE_URL}/app/views/hrm_mod/report/master_report_att_management.php"

# Visual Styles for Premium Excel
FONT_NAME = "Segoe UI"
COLOR_NAVY = "1B365D"       # Corporate Dark Navy
COLOR_SECONDARY = "2B5B84"  # Medium Slate Navy
COLOR_ACCENT = "0284C7"     # Ocean Blue
COLOR_LIGHT_BG = "F4F7FB"   # Soft Blue-Gray
COLOR_ZEBRA = "EAF5EA"      # Subtle Paste Row
COLOR_BORDER = "D1D5DB"     # Soft Gray
COLOR_PRESENT_BG = "D1FAE5" # Soft Emerald Green
COLOR_PRESENT_FG = "065F46"
COLOR_ABSENT_BG = "FEE2E2"  # Soft Red
COLOR_ABSENT_FG = "991B1B"
COLOR_LEAVE_BG = "FEF3C7"   # Soft Amber
COLOR_LEAVE_FG = "92400E"

thin_border = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_BORDER),
    bottom=Side(style='thin', color=COLOR_BORDER)
)

header_border = Border(
    left=Side(style='thin', color=COLOR_NAVY),
    right=Side(style='thin', color=COLOR_NAVY),
    top=Side(style='medium', color=COLOR_NAVY),
    bottom=Side(style='medium', color=COLOR_NAVY)
)


def format_date_for_erp(date_str):
    """Converts DD-MM-YYYY or any standard date to YYYY-MM-DD for ERP POST."""
    if not date_str:
        return "2026-09-01"
    cleaned = str(date_str).strip()
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            dt = datetime.strptime(cleaned, fmt)
            return dt.strftime("%Y-%m-%d")
        except ValueError:
            pass
    return cleaned


class MonthlyAttendanceCollector:
    def __init__(self, credentials=None, session=None):
        self.creds = credentials or get_erp_credentials()
        self.session = session or requests.Session()
        self.headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
        }
        self.session.headers.update(self.headers)
        self.is_authenticated = False
        self.records = []
        self.day_headers = []
        self.meta = {}

    def login(self):
        """Authenticates against ERP session."""
        try:
            login_data = {
                "db": self.creds.get("db", "erpcombd"),
                "cid": self.creds.get("cid", "MEP"),
                "uid": self.creds.get("uid", "10676"),
                "ibssignin": "",
                "pass": self.creds.get("pass", "96420132"),
                "submit": "Sign In"
            }
            res = self.session.post(LOGIN_URL, data=login_data, timeout=25, allow_redirects=True)
            if "login" in res.url.lower() and "home.php" not in res.url.lower():
                raise Exception("MEP ERP Login failed! Check credentials.")
            self.is_authenticated = True
            return True
        except Exception as e:
            print(f"[!] Login exception: {e}")
            raise

    def collect(self, f_date_str="01-09-2026", t_date_str="21-09-2026"):
        """
        Submits request for Report 20220522 (Monthly Attendence Sheet)
        with specified From Date and To Date.
        """
        if not self.is_authenticated:
            self.login()

        f_erp = format_date_for_erp(f_date_str)
        t_erp = format_date_for_erp(t_date_str)
        print(f"[*] Monthly Attendance Collector: Requesting range {f_erp} to {t_erp}...")

        payload = {
            "report": "20220522",
            "PBI_JOB_STATUS": "In Service",
            "fdate": f_erp,
            "tdate": t_erp,
            "submit": "SHOW"
        }

        res = self.session.post(REPORT_POST_URL, data=payload, timeout=60)
        if res.status_code != 200:
            raise Exception(f"ERP HTTP Error: Status {res.status_code}")

        soup = BeautifulSoup(res.text, "html.parser")
        tables = soup.find_all("table")
        if len(tables) < 2:
            print(f"[!] Warning: Expected at least 2 tables, found {len(tables)}.")
            if len(tables) == 1:
                data_table = tables[0]
            else:
                raise Exception("No tables found in ERP response.")
        else:
            data_table = tables[1]

        rows = data_table.find_all("tr")
        if not rows:
            raise Exception("Attendance table contains no rows.")

        header_cells = [c.text.strip().replace("\n", " ") for c in rows[0].find_all(["th", "td"])]
        # Expected: ['SL', 'Emp Id', 'EMP Name', 'Company', 'Department', 'Section', 'Unit', 'Job Location', 'September 1 Tue', ..., 'Summary']
        base_cols_count = 8
        self.day_headers = header_cells[base_cols_count:-1] if len(header_cells) > base_cols_count + 1 else []
        summary_col_present = len(header_cells) > base_cols_count and 'summary' in header_cells[-1].lower()

        parsed_records = []
        total_present_count = 0
        total_absent_count = 0
        total_leave_count = 0

        for r_idx in range(1, len(rows)):
            tds = rows[r_idx].find_all(["td", "th"])
            if not tds or len(tds) < 3:
                continue

            sl = tds[0].text.strip().replace("\xa0", " ") if len(tds) > 0 else str(r_idx)
            emp_id = tds[1].text.strip().replace("\xa0", " ") if len(tds) > 1 else ""
            emp_name = tds[2].text.strip().replace("\xa0", " ") if len(tds) > 2 else ""
            company = tds[3].text.strip().replace("\xa0", " ") if len(tds) > 3 else ""
            department = tds[4].text.strip().replace("\xa0", " ") if len(tds) > 4 else ""
            section = tds[5].text.strip().replace("\xa0", " ") if len(tds) > 5 else ""
            unit = tds[6].text.strip().replace("\xa0", " ") if len(tds) > 6 else ""
            job_location = tds[7].text.strip().replace("\xa0", " ") if len(tds) > 7 else ""

            daily_records = []
            row_present = 0
            row_absent = 0
            row_leave = 0

            # Daily cells
            day_tds = tds[base_cols_count:base_cols_count + len(self.day_headers)]
            for d_idx, td in enumerate(day_tds):
                d_head = self.day_headers[d_idx] if d_idx < len(self.day_headers) else f"Day {d_idx+1}"
                day_txt = td.get_text(separator="\n").strip().replace("\xa0", " ")

                times = re.findall(r'(\d{1,2}:\d{2}\s*(?:am|pm)?)', day_txt, re.IGNORECASE)
                in_time = times[0] if len(times) > 0 else ""
                out_time = times[1] if len(times) > 1 else ""

                day_lower = day_txt.lower()
                if "present" in day_lower:
                    status_clean = "Present"
                elif "offday" in day_lower or "holiday" in day_lower or "weekly" in day_lower:
                    status_clean = "OFFDAY"
                elif "casual" in day_lower:
                    status_clean = "Casual Leave"
                elif "leave" in day_lower:
                    status_clean = "Leave"
                elif "late" in day_lower:
                    status_clean = "Late"
                elif "early" in day_lower:
                    status_clean = "EarlyOut"
                elif day_txt.strip().upper() in ["A", "ABSENT"]:
                    status_clean = "Absent"
                else:
                    status_clean = day_txt.split("\n")[0].strip() if day_txt else "-"

                if status_clean == "Present":
                    row_present += 1
                    total_present_count += 1
                elif status_clean == "Absent":
                    row_absent += 1
                    total_absent_count += 1
                elif "leave" in status_clean.lower():
                    row_leave += 1
                    total_leave_count += 1

                daily_records.append({
                    "day_header": d_head,
                    "status": status_clean,
                    "status_type": status_clean,
                    "in_time": in_time,
                    "out_time": out_time,
                    "details": day_txt
                })

            summary_text = tds[-1].text.strip().replace("\xa0", " ") if summary_col_present and len(tds) > len(self.day_headers) + base_cols_count else f"P:{row_present} A:{row_absent} L:{row_leave}"

            rec = {
                "sl": sl,
                "emp_id": emp_id,
                "emp_name": emp_name,
                "company": company,
                "department": department,
                "section": section,
                "unit": unit,
                "job_location": job_location,
                "present_days": row_present,
                "absent_days": row_absent,
                "leave_days": row_leave,
                "daily": daily_records,
                "summary": summary_text
            }
            parsed_records.append(rec)

        self.records = parsed_records
        self.meta = {
            "report_name": "Monthly Attendence Sheet",
            "report_id": "20220522",
            "source_path": "HRM Module ⟶ New-HRM Report ⟶ Reports ⟶ Monthly Attendence Sheet ⟶ SHOW",
            "from_date": f_date_str,
            "to_date": t_date_str,
            "fdate_erp": f_erp,
            "tdate_erp": t_erp,
            "total_employees": len(parsed_records),
            "total_presents": total_present_count,
            "total_absents": total_absent_count,
            "total_leaves": total_leave_count,
            "days_count": len(self.day_headers),
            "collected_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }

        print(f"[+] Successfully extracted {len(parsed_records)} employee attendance records across {len(self.day_headers)} days.")
        return self.records

    def save_cache(self, path=None):
        """Saves collected data to JSON cache with Sync Failure Protection."""
        target_path = path or CACHE_FILE
        if not self.records and os.path.exists(target_path):
            print("[!] Empty records returned. Preserving existing valid cache!")
            return False

        payload = {
            "meta": self.meta,
            "day_headers": self.day_headers,
            "items": self.records,
            "summary": {
                "total_records": len(self.records),
                "total_presents": self.meta.get("total_presents", 0),
                "total_absents": self.meta.get("total_absents", 0),
                "total_leaves": self.meta.get("total_leaves", 0),
                "last_sync": datetime.now().strftime("%Y-%m-%d %I:%M %p")
            }
        }
        with open(target_path, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2, ensure_ascii=False)
        print(f"[+] Saved Monthly Attendance cache to: {target_path}")
        return True

    def export_to_excel(self, output_path=None):
        """Generates an executive-styled multi-column Excel workbook matching ERP exact structure."""
        out_file = output_path or EXCEL_OUTPUT_FILE
        if not self.records:
            return False

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Monthly Attendance Sheet"
        ws.views.sheetView[0].showGridLines = True

        # 1. Header Banner
        ws.merge_cells("A1:K1")
        title_cell = ws["A1"]
        title_cell.value = "MEP GROUP — MONTHLY ATTENDANCE SHEET"
        title_cell.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        # Subtitle
        ws.merge_cells("A2:K2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Period: {self.meta.get('from_date')} to {self.meta.get('to_date')} | Total Staff: {len(self.records)} | Generated: {self.meta.get('collected_at')}"
        sub_cell.font = Font(name=FONT_NAME, size=10, italic=True, color="FFFFFF")
        sub_cell.fill = PatternFill(start_color=COLOR_SECONDARY, end_color=COLOR_SECONDARY, fill_type="solid")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        # 2. Table Headers
        headers = ["SL", "Emp ID", "Employee Name", "Company", "Department", "Section", "Unit", "Job Location"]
        headers.extend(self.day_headers)
        headers.extend(["Present", "Absent", "Leave", "Summary"])

        ws.append([]) # Row 3 blank
        ws.append(headers) # Row 4
        header_row_idx = 4
        ws.row_dimensions[header_row_idx].height = 28

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = header_border

        # 3. Data Rows
        current_row = 5
        for it in self.records:
            row_data = [
                it.get("sl", ""),
                it.get("emp_id", ""),
                it.get("emp_name", ""),
                it.get("company", ""),
                it.get("department", ""),
                it.get("section", ""),
                it.get("unit", ""),
                it.get("job_location", "")
            ]
            for d in it.get("daily", []):
                row_data.append(d.get("details", ""))
            row_data.extend([
                it.get("present_days", 0),
                it.get("absent_days", 0),
                it.get("leave_days", 0),
                it.get("summary", "")
            ])

            ws.append(row_data)
            ws.row_dimensions[current_row].height = 20
            fill_color = COLOR_ZEBRA if (current_row % 2 == 0) else "FFFFFF"

            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = Font(name=FONT_NAME, size=9)
                cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", horizontal="left" if col_idx in [3, 4, 5, 6, 7] else "center")

            current_row += 1

        # Auto column dimensions
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 30), 10)

        wb.save(out_file)
        print(f"[+] Saved Monthly Attendance Excel to: {out_file}")
        return True


if __name__ == "__main__":
    collector = MonthlyAttendanceCollector()
    collector.login()
    collector.collect("01-09-2026", "21-09-2026")
    collector.save_cache()
    collector.export_to_excel()
