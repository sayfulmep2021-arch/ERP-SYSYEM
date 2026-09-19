"""
MEP Group ERP - Enterprise Reports Collector
Automated Collectors for:
  1. Inventory Movement Report (221023) - Fan Assemble (warehouse_id: 86)
  2. Inventory Movement Report (221023) - Armature Winding (warehouse_id: 83)
  3. Inventory Movement Report (221023) - Semi Finished Goods (item_group: 1300000000)
  4. Stock Position Report Detail Closing (91223) - Ground Floor Fan Store FAN-1 (warehouse_id: 113)
  5. Stock Movement Report (222) - Finished Goods at FAN Floor (item_sub_group: 100100000, warehouse_id: 68)

Date Range: 1st of Current Month to Running Date (Today)
Author: Antigravity
"""

import os
import json
import time
from datetime import datetime
import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_URL = "https://www.mepgrouperp.com/1027"
LOGIN_URL = f"{BASE_URL}/login/pages/main/index.php"
PROD_REPORT_URL = f"{BASE_URL}/production_mod/pages/report/master_report.php"
WH_REPORT_URL = f"{BASE_URL}/warehouse_mod/pages/report/master_report.php"

from erp_credentials import get_erp_credentials

DEFAULT_CREDENTIALS = get_erp_credentials()

# Cache & Excel Output File Paths
FAN_ASSEMBLE_CACHE = "fan_assemble_cache.json"
FAN_ASSEMBLE_EXCEL = "Fan_Assemble_Inventory_Movement.xlsx"

ARMATURE_CACHE = "armature_winding_cache.json"
ARMATURE_EXCEL = "Armature_Winding_Inventory_Movement.xlsx"

SEMI_FINISHED_CACHE = "semi_finished_cache.json"
SEMI_FINISHED_EXCEL = "Semi_Finished_Inventory_Movement.xlsx"

FAN_STORE_CACHE = "fan_store_stock_cache.json"
FAN_STORE_EXCEL = "Ground_Floor_Fan_Store_Stock_Detail.xlsx"

STOCK_MOVEMENT_CACHE = "stock_movement_cache.json"
STOCK_MOVEMENT_EXCEL = "Finished_Goods_FAN_Floor_Stock_Movement.xlsx"

FONT_NAME = "Segoe UI"
COLOR_NAVY = "1B365D"
COLOR_SECONDARY = "2B5B84"
COLOR_LIGHT_BG = "F4F7FB"
COLOR_ZEBRA = "EAF5EA"  # Executive Paste Color
COLOR_BORDER = "D1D5DB"

thin_border = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_BORDER),
    bottom=Side(style='thin', color=COLOR_BORDER)
)

header_border = Border(
    left=Side(style='thin', color=COLOR_NAVY),
    right=Side(style='thin', color=COLOR_NAVY),
    top=Side(style='medium', color=COLOR_NAVY),
    bottom=Side(style='medium', color=COLOR_NAVY)
)


def get_default_date_range():
    """Returns (1st of current month, today) in YYYY-MM-DD format."""
    now = datetime.now()
    f_date = datetime(now.year, now.month, 1).strftime("%Y-%m-%d")
    t_date = now.strftime("%Y-%m-%d")
    return f_date, t_date


class ERPBaseCollector:
    def __init__(self, credentials=None, session=None):
        self.creds = credentials or DEFAULT_CREDENTIALS
        if session:
            self.session = session
            self._shared_session = True
        else:
            self.session = requests.Session()
            self.session.headers.update({
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                "Accept-Language": "en-US,en;q=0.9",
            })
            self._shared_session = False

    def login(self):
        if getattr(self, '_shared_session', False):
            return True
        print("[*] Logging into MEP Group ERP...")
        res = self.session.post(LOGIN_URL, data=self.creds, timeout=30)
        if "login" in res.url.lower() and "home.php" not in res.url.lower():
            raise Exception("Authentication failed! Please verify credentials.")
        print("[+] Logged in successfully!")
        return True


# =============================================================================
# 1, 2, 3: INVENTORY MOVEMENT REPORT (221023)
# =============================================================================
class InventoryMovementCollector(ERPBaseCollector):
    """
    Handles Production Module -> Advance Production Reports -> Inventory Movement Report (221023)
    Subtypes:
      - 'fan_assemble': warehouse_id = 86
      - 'armature_winding': warehouse_id = 83
      - 'semi_finished': item_group = 1300000000
    """
    CONFIGS = {
        "fan_assemble": {
            "title": "Inventory Movement — Fan Assemble",
            "short_title": "Fan Assemble",
            "report_code": "221023",
            "payload_extra": {"warehouse_id": "86"},
            "cache_file": FAN_ASSEMBLE_CACHE,
            "excel_file": FAN_ASSEMBLE_EXCEL,
            "source": "Production Module ⟶ Advance Production Reports ⟶ Inventory Movement Report(221023) ⟶ Inventory Name: Fan Assemble"
        },
        "armature_winding": {
            "title": "Inventory Movement — Armature Winding",
            "short_title": "Armature Winding",
            "report_code": "221023",
            "payload_extra": {"warehouse_id": "83"},
            "cache_file": ARMATURE_CACHE,
            "excel_file": ARMATURE_EXCEL,
            "source": "Production Module ⟶ Advance Production Reports ⟶ Inventory Movement Report(221023) ⟶ Inventory Name: Armature Winding"
        },
        "semi_finished": {
            "title": "Inventory Movement — Semi Finished Goods",
            "short_title": "Semi Finished Goods",
            "report_code": "221023",
            "payload_extra": {"item_group": "1300000000"},
            "cache_file": SEMI_FINISHED_CACHE,
            "excel_file": SEMI_FINISHED_EXCEL,
            "source": "Production Module ⟶ Advance Production Reports ⟶ Inventory Movement Report(221023) ⟶ Item Group: Semi Finished Goods"
        }
    }

    def __init__(self, sub_type="fan_assemble", credentials=None, session=None):
        super().__init__(credentials, session=session)
        if sub_type not in self.CONFIGS:
            raise ValueError(f"Unknown subtype: {sub_type}. Must be one of {list(self.CONFIGS.keys())}")
        self.sub_type = sub_type
        self.cfg = self.CONFIGS[sub_type]
        self.items = []
        self.meta = {}
        self.raw_headers = []
        self.raw_rows = []
        self.total_row = []
        self.report_title = ""
        self.report_subtitle = ""

    def collect(self, f_date_str=None, t_date_str=None):
        def_f, def_t = get_default_date_range()
        f_date = f_date_str or def_f
        t_date = t_date_str or def_t

        print(f"[*] Querying {self.cfg['title']} from {f_date} to {t_date}...")
        payload = {
            "report": self.cfg["report_code"],
            "group_for": "3",  # Company: FAN
            "f_date": f_date,
            "t_date": t_date,
            "submit": "Report"
        }
        payload.update(self.cfg["payload_extra"])

        res = self.session.post(PROD_REPORT_URL, data=payload, timeout=60)
        soup = BeautifulSoup(res.text, "html.parser")
        tbl = soup.find("table")
        if not tbl:
            raise Exception(f"Report table not found in ERP response for {self.cfg['title']}.")

        h2 = soup.find("h2")
        h3 = soup.find("h3")
        self.report_title = h2.get_text(" ", strip=True) if h2 else "Stock Movement Report"
        self.report_subtitle = h3.get_text(" ", strip=True) if h3 else f"Data Interval: {f_date} to {t_date}"

        rows = tbl.find_all("tr")
        parsed = []
        raw_rows = []

        # Row 0 is header; capture exact header attributes
        if len(rows) > 0:
            h_cells = []
            for th in rows[0].find_all(["th", "td"]):
                h_cells.append({
                    "text": th.get_text(" ", strip=True),
                    "bgcolor": th.get("bgcolor", "") or "",
                    "colspan": int(th.get("colspan", 1)),
                    "rowspan": int(th.get("rowspan", 1)),
                    "style": th.get("style", "") or ""
                })
            self.raw_headers = [h_cells]

        # Row 0 is header; data starts from row 1
        for r in rows[1:]:
            cols = [c.get_text(" ", strip=True) for c in r.find_all(["th", "td"])]
            if cols and cols[0].isdigit():
                raw_rows.append(cols)
                try:
                    opening = float(cols[7].replace(",", "") or 0)
                    opening_adjust = float(cols[8].replace(",", "") or 0)
                    store_receive = float(cols[9].replace(",", "") or 0)
                    section_receive = float(cols[10].replace(",", "") or 0)
                    production_receive = float(cols[11].replace(",", "") or 0)
                    damage_receive = float(cols[12].replace(",", "") or 0)
                    others_receive = float(cols[13].replace(",", "") or 0)
                    total_receive = float(cols[14].replace(",", "") or 0)
                    consumption = float(cols[15].replace(",", "") or 0)
                    wip_issue = float(cols[16].replace(",", "") or 0)
                    issue_opening_adjust = float(cols[17].replace(",", "") or 0)
                    issue_to_repair = float(cols[18].replace(",", "") or 0)
                    issue_to_damage = float(cols[19].replace(",", "") or 0)
                    in_transit = float(cols[20].replace(",", "") or 0)
                    total_issue = float(cols[21].replace(",", "") or 0)
                    closing = float(cols[22].replace(",", "") or 0)
                    bin_closing = float(cols[23].replace(",", "") or 0)
                    diff = float(cols[24].replace(",", "") or 0)
                except (ValueError, IndexError):
                    continue

                parsed.append({
                    "sl": cols[0],
                    "company": cols[1],
                    "category": cols[2],
                    "erp_code": cols[3],
                    "code": cols[4],
                    "item_name": cols[5],
                    "unit": cols[6] or "Pcs",
                    "opening": opening,
                    "opening_adjust": opening_adjust,
                    "store_receive": store_receive,
                    "section_receive": section_receive,
                    "production_receive": production_receive,
                    "damage_receive": damage_receive,
                    "others_receive": others_receive,
                    "total_receive": total_receive,
                    "consumption": consumption,
                    "wip_issue": wip_issue,
                    "issue_opening_adjust": issue_opening_adjust,
                    "issue_to_repair": issue_to_repair,
                    "issue_to_damage": issue_to_damage,
                    "in_transit": in_transit,
                    "total_issue": total_issue,
                    "closing": closing,
                    "bin_closing": bin_closing,
                    "diff": diff,
                    "remarks": cols[25] if len(cols) > 25 else ""
                })

        self.raw_rows = raw_rows

        # Capture total summary row from the bottom if present
        self.total_row = []
        if len(rows) > 1:
            last_cols = [c.get_text(" ", strip=True) for c in rows[-1].find_all(["th", "td"])]
            if any(last_cols) and not last_cols[0].isdigit():
                self.total_row = last_cols

        self.items = parsed
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.meta = {
            "title": self.cfg["title"],
            "short_title": self.cfg["short_title"],
            "sub_type": self.sub_type,
            "source": self.cfg["source"],
            "company": "FAN",
            "report_code": self.cfg["report_code"],
            "f_date": f_date,
            "t_date": t_date,
            "collected_at": now_str,
            "total_items": len(parsed),
            "active_items": sum(1 for it in parsed if (it["total_receive"] > 0 or it["total_issue"] > 0 or it["closing"] > 0 or it["opening"] > 0)),
            "total_opening": sum(it["opening"] for it in parsed),
            "total_receive": sum(it["total_receive"] for it in parsed),
            "total_issue": sum(it["total_issue"] for it in parsed),
            "total_closing": sum(it["closing"] for it in parsed),
            "total_production_receive": sum(it["production_receive"] for it in parsed),
            "total_consumption": sum(it["consumption"] for it in parsed)
        }

        print(f"[+] Collected {len(parsed)} items for {self.cfg['title']} (Closing Stock: {self.meta['total_closing']:,.2f})")
        return self.items

    def save_cache(self):
        filepath = self.cfg["cache_file"]
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump({
                "meta": self.meta,
                "items": self.items,
                "raw_headers": self.raw_headers,
                "raw_rows": self.raw_rows,
                "total_row": self.total_row,
                "report_title": self.report_title,
                "report_subtitle": self.report_subtitle
            }, f, indent=2, ensure_ascii=False)
        print(f"[+] Saved cache to {filepath}")

    def load_cache(self):
        filepath = self.cfg["cache_file"]
        if os.path.exists(filepath):
            with open(filepath, "r", encoding="utf-8") as f:
                d = json.load(f)
                self.items = d.get("items", [])
                self.meta = d.get("meta", {})
                self.raw_headers = d.get("raw_headers", [])
                self.raw_rows = d.get("raw_rows", [])
                self.total_row = d.get("total_row", [])
                self.report_title = d.get("report_title", "")
                self.report_subtitle = d.get("report_subtitle", "")
                return True
        return False

    def export_to_excel(self, output_file=None):
        out_file = output_file or self.cfg["excel_file"]
        print(f"[*] Exporting {self.cfg['title']} Workbook to {out_file}...")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        # 1. Executive Summary Sheet
        ws_sum = wb.create_sheet(title="Executive Summary")
        ws_sum.views.sheetView[0].showGridLines = True
        self._build_summary_sheet(ws_sum)

        # 2. Master Table Sheet
        ws_main = wb.create_sheet(title="Inventory Movement")
        ws_main.views.sheetView[0].showGridLines = True
        self._build_table_sheet(ws_main)

        wb.save(out_file)
        print(f"[+] Successfully generated {out_file}!")
        return out_file

    def _build_summary_sheet(self, ws):
        ws.merge_cells("A1:H2")
        b = ws["A1"]
        b.value = f"MEP GROUP  |  FAN — {self.cfg['title']}"
        b.font = Font(name=FONT_NAME, size=15, bold=True, color="FFFFFF")
        b.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        b.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        f_d = self.meta.get("f_date", "")
        t_d = self.meta.get("t_date", "")
        ws["A3"].value = f"{self.cfg['source']}  |  Company: FAN  |  Date Range: {f_d} to {t_d}"
        ws["A3"].font = Font(name=FONT_NAME, size=9.5, italic=True, color="64748B")

        kpis = [
            ("Total Catalog Items", len(self.items), "Inventory Items in Scope", "1B365D"),
            ("Active Movement Items", self.meta.get("active_items", 0), "Items with In/Out Activity", "059669"),
            ("Total Received", f"{self.meta.get('total_receive', 0):,.2f}", "Cumulative Inflow", "0284C7"),
            ("Total Closing Stock", f"{self.meta.get('total_closing', 0):,.2f}", "Current Net Closing Units", "D97706"),
        ]

        col_pairs = [("B", "C"), ("D", "E"), ("F", "G"), ("H", "I")]
        for idx, (lbl, val, sub, col_hex) in enumerate(kpis):
            c1, c2 = col_pairs[idx]
            ws.merge_cells(f"{c1}5:{c2}5")
            ws.merge_cells(f"{c1}6:{c2}6")
            ws.merge_cells(f"{c1}7:{c2}7")

            v_cell = ws[f"{c1}5"]
            v_cell.value = val
            v_cell.font = Font(name=FONT_NAME, size=16, bold=True, color=col_hex)
            v_cell.alignment = Alignment(horizontal="center", vertical="center")

            l_cell = ws[f"{c1}6"]
            l_cell.value = lbl
            l_cell.font = Font(name=FONT_NAME, size=10, bold=True, color="334155")
            l_cell.alignment = Alignment(horizontal="center", vertical="center")

            s_cell = ws[f"{c1}7"]
            s_cell.value = sub
            s_cell.font = Font(name=FONT_NAME, size=8, color="64748B")
            s_cell.alignment = Alignment(horizontal="center", vertical="center")

            for r in range(5, 8):
                for c in [c1, c2]:
                    ws[f"{c}{r}"].fill = PatternFill(start_color=COLOR_LIGHT_BG, end_color=COLOR_LIGHT_BG, fill_type="solid")
                    ws[f"{c}{r}"].border = thin_border

        for col in ["B", "C", "D", "E", "F", "G", "H", "I"]:
            ws.column_dimensions[col].width = 18

    def _build_table_sheet(self, ws):
        ws.merge_cells("A1:T1")
        title = ws["A1"]
        title.value = f"MEP GROUP  |  FAN — {self.cfg['title']} (Report {self.cfg['report_code']})"
        title.font = Font(name=FONT_NAME, size=13, bold=True, color="FFFFFF")
        title.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 28

        headers = [
            ("SL", 7, "center"),
            ("Company", 10, "center"),
            ("Category", 18, "left"),
            ("ERP Code", 14, "center"),
            ("Code", 18, "left"),
            ("Item Name", 36, "left"),
            ("Unit", 8, "center"),
            ("Opening", 14, "right"),
            ("Store Rec.", 13, "right"),
            ("Section Rec.", 13, "right"),
            ("Prod Rec.", 13, "right"),
            ("Total Rec.", 15, "right"),
            ("Consumption", 14, "right"),
            ("WIP Issue", 14, "right"),
            ("Total Issue", 15, "right"),
            ("Closing", 15, "right"),
            ("Bin Closing", 14, "right"),
            ("Diff", 12, "right"),
            ("Remarks", 16, "left")
        ]

        ws.row_dimensions[3].height = 24
        for c_idx, (h_title, _, align) in enumerate(headers, start=1):
            c = ws.cell(row=3, column=c_idx, value=h_title)
            c.font = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            c.alignment = Alignment(horizontal=align, vertical="center")
            c.border = header_border

        curr_r = 4
        for it in self.items:
            bg = COLOR_ZEBRA if curr_r % 2 == 1 else "FFFFFF"
            ws.row_dimensions[curr_r].height = 20
            row_vals = [
                it["sl"],
                it["company"],
                it["category"],
                it["erp_code"],
                it["code"],
                it["item_name"],
                it["unit"],
                it["opening"],
                it["store_receive"],
                it["section_receive"],
                it["production_receive"],
                it["total_receive"],
                it["consumption"],
                it["wip_issue"],
                it["total_issue"],
                it["closing"],
                it["bin_closing"],
                it["diff"],
                it["remarks"]
            ]
            for c_idx, v in enumerate(row_vals, start=1):
                c = ws.cell(row=curr_r, column=c_idx, value=v)
                c.font = Font(name=FONT_NAME, size=9)
                c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                c.border = thin_border
                c.alignment = Alignment(horizontal=headers[c_idx-1][2], vertical="center")
                if c_idx in [8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18]:
                    c.number_format = "#,##0.00"
            curr_r += 1

        # Summary row
        ws.row_dimensions[curr_r].height = 24
        ws.cell(row=curr_r, column=1, value="Total")
        ws.cell(row=curr_r, column=6, value=f"{len(self.items)} Items")
        ws.cell(row=curr_r, column=8, value=self.meta.get("total_opening", 0))
        ws.cell(row=curr_r, column=12, value=self.meta.get("total_receive", 0))
        ws.cell(row=curr_r, column=15, value=self.meta.get("total_issue", 0))
        ws.cell(row=curr_r, column=16, value=self.meta.get("total_closing", 0))

        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(row=curr_r, column=c_idx)
            c.font = Font(name=FONT_NAME, size=10, bold=True, color="000000")
            c.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            c.border = Border(top=Side(style='thin'), bottom=Side(style='double'))
            if c_idx in [8, 12, 15, 16]:
                c.number_format = "#,##0.00"

        for c_idx, (_, width, _) in enumerate(headers, start=1):
            col_letter = get_column_letter(c_idx)
            ws.column_dimensions[col_letter].width = width


# =============================================================================
# 4: STOCK POSITION REPORT DETAIL (CLOSING) (91223) - GROUND FLOOR FAN STORE
# =============================================================================
class FanStoreStockCollector(ERPBaseCollector):
    """
    Warehouse Module -> Warehouse Reports -> Stock Position Report Detail(Closing)(91223)
    Company: FAN (group_for = 3)
    Warehouse Name: Ground Floor Fan Store (FAN-1) (warehouse_id = 113)
    """
    def __init__(self, credentials=None, session=None):
        super().__init__(credentials, session=session)
        self.items = []
        self.meta = {}
        self.raw_headers = []
        self.raw_rows = []
        self.total_row = []
        self.report_title = ""
        self.report_subtitle = ""

    def collect(self, f_date_str=None, t_date_str=None):
        def_f, def_t = get_default_date_range()
        f_date = f_date_str or def_f
        t_date = t_date_str or def_t

        print(f"[*] Querying Stock Position Detail (91223) for Ground Floor Fan Store (FAN-1) from {f_date} to {t_date}...")
        payload = {
            "report": "91223",
            "group_for": "3",  # FAN
            "warehouse_id": "113",  # Ground Floor Fan Store (FAN-1)
            "f_date": f_date,
            "t_date": t_date,
            "submit": "Report"
        }

        res = self.session.post(WH_REPORT_URL, data=payload, timeout=60)
        soup = BeautifulSoup(res.text, "html.parser")
        tbl = soup.find("table")
        if not tbl:
            raise Exception("Stock Position Detail table not found in ERP response.")

        self.report_title = "Warehouse Stock Position Report Detail(Closing)"
        self.report_subtitle = f"Closing Stock of Date-{t_date}"

        rows = tbl.find_all("tr")
        parsed = []
        raw_rows = []

        # Row 0: Banner, Row 1: Header (11 columns)
        if len(rows) > 1:
            h_cells = []
            for cell in rows[1].find_all(["th", "td"]):
                h_cells.append({
                    "text": cell.get_text(" ", strip=True),
                    "bgcolor": cell.get("bgcolor", "") or "",
                    "colspan": int(cell.get("colspan", 1)),
                    "rowspan": int(cell.get("rowspan", 1)),
                    "style": cell.get("style", "") or ""
                })
            self.raw_headers = [h_cells]

        # Data starts at Row 2
        for r in rows[2:]:
            cols = [c.get_text(" ", strip=True) for c in r.find_all(["th", "td"])]
            if cols and cols[0].isdigit():
                raw_rows.append(cols)
                try:
                    store_qty = float(cols[9].replace(",", "") or 0)
                    section_qty = float(cols[10].replace(",", "") or 0)
                    total_qty = store_qty + section_qty
                except (ValueError, IndexError):
                    continue

                parsed.append({
                    "sl": cols[0],
                    "company": cols[1],
                    "group": cols[2],
                    "category": cols[3],
                    "subcategory": cols[4],
                    "item_code": cols[5],
                    "fg": cols[6],
                    "item_name": cols[7],
                    "unit": cols[8] or "Pcs",
                    "store_qty": store_qty,
                    "section_qty": section_qty,
                    "total_qty": total_qty
                })

        self.raw_rows = raw_rows

        # Capture total summary row from the bottom if present
        self.total_row = []
        if len(rows) > 2:
            last_cols = [c.get_text(" ", strip=True) for c in rows[-1].find_all(["th", "td"])]
            if any(last_cols) and not last_cols[0].isdigit():
                self.total_row = last_cols

        self.items = parsed
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        tot_store = sum(it["store_qty"] for it in parsed)
        tot_sec = sum(it["section_qty"] for it in parsed)
        tot_qty = tot_store + tot_sec

        self.meta = {
            "title": "Stock Position Report Detail (Closing) — Ground Floor Fan Store (FAN-1)",
            "short_title": "Ground Floor Fan Store (FAN-1)",
            "source": "Warehouse Module ⟶ Warehouse Reports ⟶ Stock Position Report Detail(Closing)(91223) ⟶ Warehouse: Ground Floor Fan Store (FAN-1)",
            "company": "FAN",
            "report_code": "91223",
            "warehouse_name": "Ground Floor Fan Store (FAN-1)",
            "f_date": f_date,
            "t_date": t_date,
            "collected_at": now_str,
            "total_items": len(parsed),
            "in_stock_items": sum(1 for it in parsed if it["total_qty"] > 0),
            "zero_stock_items": sum(1 for it in parsed if it["total_qty"] <= 0),
            "total_store_qty": tot_store,
            "total_section_qty": tot_sec,
            "total_quantity": tot_qty
        }

        print(f"[+] Collected {len(parsed)} items for Ground Floor Fan Store (FAN-1) (Total Qty: {tot_qty:,.2f})")
        return self.items

    def save_cache(self, filepath=FAN_STORE_CACHE):
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump({
                "meta": self.meta,
                "items": self.items,
                "raw_headers": self.raw_headers,
                "raw_rows": self.raw_rows,
                "total_row": self.total_row,
                "report_title": self.report_title,
                "report_subtitle": self.report_subtitle
            }, f, indent=2, ensure_ascii=False)
        print(f"[+] Saved cache to {filepath}")

    def load_cache(self, filepath=FAN_STORE_CACHE):
        if os.path.exists(filepath):
            with open(filepath, "r", encoding="utf-8") as f:
                d = json.load(f)
                self.items = d.get("items", [])
                self.meta = d.get("meta", {})
                self.raw_headers = d.get("raw_headers", [])
                self.raw_rows = d.get("raw_rows", [])
                self.total_row = d.get("total_row", [])
                self.report_title = d.get("report_title", "")
                self.report_subtitle = d.get("report_subtitle", "")
                return True
        return False

    def export_to_excel(self, output_file=FAN_STORE_EXCEL):
        print(f"[*] Exporting Ground Floor Fan Store Workbook to {output_file}...")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        # 1. Executive Summary Sheet
        ws_sum = wb.create_sheet(title="Executive Summary")
        ws_sum.views.sheetView[0].showGridLines = True
        self._build_summary_sheet(ws_sum)

        # 2. Master Stock Detail
        ws_main = wb.create_sheet(title="Stock Position Detail")
        ws_main.views.sheetView[0].showGridLines = True
        self._build_table_sheet(ws_main)

        wb.save(output_file)
        print(f"[+] Successfully generated {output_file}!")
        return output_file

    def _build_summary_sheet(self, ws):
        ws.merge_cells("A1:H2")
        b = ws["A1"]
        b.value = "MEP GROUP  |  FAN Ground Floor Fan Store (FAN-1) — Stock Position Detail (91223)"
        b.font = Font(name=FONT_NAME, size=15, bold=True, color="FFFFFF")
        b.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        b.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        f_d = self.meta.get("f_date", "")
        t_d = self.meta.get("t_date", "")
        ws["A3"].value = f"{self.meta.get('source', '')}  |  Date Range: {f_d} to {t_d}"
        ws["A3"].font = Font(name=FONT_NAME, size=9.5, italic=True, color="64748B")

        kpis = [
            ("Total Catalog Items", len(self.items), "Registered Inventory SKUs", "1B365D"),
            ("Active In-Stock Items", self.meta.get("in_stock_items", 0), "Items with Total Qty > 0", "059669"),
            ("Store Quantity", f"{self.meta.get('total_store_qty', 0):,.2f}", "Main Ground Store Stock", "0284C7"),
            ("Total Warehouse Qty", f"{self.meta.get('total_quantity', 0):,.2f}", "Combined Store & Section Stock", "D97706"),
        ]

        col_pairs = [("B", "C"), ("D", "E"), ("F", "G"), ("H", "I")]
        for idx, (lbl, val, sub, col_hex) in enumerate(kpis):
            c1, c2 = col_pairs[idx]
            ws.merge_cells(f"{c1}5:{c2}5")
            ws.merge_cells(f"{c1}6:{c2}6")
            ws.merge_cells(f"{c1}7:{c2}7")

            v_cell = ws[f"{c1}5"]
            v_cell.value = val
            v_cell.font = Font(name=FONT_NAME, size=16, bold=True, color=col_hex)
            v_cell.alignment = Alignment(horizontal="center", vertical="center")

            l_cell = ws[f"{c1}6"]
            l_cell.value = lbl
            l_cell.font = Font(name=FONT_NAME, size=10, bold=True, color="334155")
            l_cell.alignment = Alignment(horizontal="center", vertical="center")

            s_cell = ws[f"{c1}7"]
            s_cell.value = sub
            s_cell.font = Font(name=FONT_NAME, size=8, color="64748B")
            s_cell.alignment = Alignment(horizontal="center", vertical="center")

            for r in range(5, 8):
                for c in [c1, c2]:
                    ws[f"{c}{r}"].fill = PatternFill(start_color=COLOR_LIGHT_BG, end_color=COLOR_LIGHT_BG, fill_type="solid")
                    ws[f"{c}{r}"].border = thin_border

        for col in ["B", "C", "D", "E", "F", "G", "H", "I"]:
            ws.column_dimensions[col].width = 18

    def _build_table_sheet(self, ws):
        ws.merge_cells("A1:K1")
        title = ws["A1"]
        title.value = "MEP GROUP  |  FAN Ground Floor Fan Store (FAN-1) — Stock Position Detail (91223)"
        title.font = Font(name=FONT_NAME, size=13, bold=True, color="FFFFFF")
        title.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 28

        headers = [
            ("SL", 7, "center"),
            ("Company", 10, "center"),
            ("SubCategory", 22, "left"),
            ("Item Code", 15, "center"),
            ("FG Code", 15, "center"),
            ("Item Name & Description", 38, "left"),
            ("Unit", 8, "center"),
            ("Store Qty", 16, "right"),
            ("Section Qty", 16, "right"),
            ("Total Quantity", 18, "right")
        ]

        ws.row_dimensions[3].height = 24
        for c_idx, (h_title, _, align) in enumerate(headers, start=1):
            c = ws.cell(row=3, column=c_idx, value=h_title)
            c.font = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            c.alignment = Alignment(horizontal=align, vertical="center")
            c.border = header_border

        curr_r = 4
        for it in self.items:
            bg = COLOR_ZEBRA if curr_r % 2 == 1 else "FFFFFF"
            ws.row_dimensions[curr_r].height = 20
            row_vals = [
                it["sl"],
                it["company"],
                it["subcategory"] or it["category"] or "General",
                it["item_code"],
                it["fg"],
                it["item_name"],
                it["unit"],
                it["store_qty"],
                it["section_qty"],
                it["total_qty"]
            ]
            for c_idx, v in enumerate(row_vals, start=1):
                c = ws.cell(row=curr_r, column=c_idx, value=v)
                c.font = Font(name=FONT_NAME, size=9)
                c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                c.border = thin_border
                c.alignment = Alignment(horizontal=headers[c_idx-1][2], vertical="center")
                if c_idx in [8, 9, 10]:
                    c.number_format = "#,##0.00"
            curr_r += 1

        # Summary Row
        ws.row_dimensions[curr_r].height = 24
        ws.cell(row=curr_r, column=1, value="Total")
        ws.cell(row=curr_r, column=6, value=f"{len(self.items)} Items")
        ws.cell(row=curr_r, column=8, value=self.meta.get("total_store_qty", 0))
        ws.cell(row=curr_r, column=9, value=self.meta.get("total_section_qty", 0))
        ws.cell(row=curr_r, column=10, value=self.meta.get("total_quantity", 0))

        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(row=curr_r, column=c_idx)
            c.font = Font(name=FONT_NAME, size=10, bold=True, color="000000")
            c.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            c.border = Border(top=Side(style='thin'), bottom=Side(style='double'))
            if c_idx in [8, 9, 10]:
                c.number_format = "#,##0.00"

        for c_idx, (_, width, _) in enumerate(headers, start=1):
            col_letter = get_column_letter(c_idx)
            ws.column_dimensions[col_letter].width = width


# =============================================================================
# 5: STOCK MOVEMENT REPORT (222) - FINISHED GOODS AT FAN FLOOR
# =============================================================================
class StockMovementCollector(ERPBaseCollector):
    """
    Warehouse Module -> Warehouse Reports -> Stock Movement Report(222)
    Company: FAN (group_for = 3)
    Item Sub Group: Finished Goods (item_sub_group = 100100000)
    Warehouse Name: FAN Floor (warehouse_id = 68)
    """
    def __init__(self, credentials=None, session=None):
        super().__init__(credentials, session=session)
        self.items = []
        self.meta = {}
        self.raw_headers = []
        self.raw_rows = []
        self.total_row = []
        self.report_title = ""
        self.report_subtitle = ""

    def collect(self, f_date_str=None, t_date_str=None):
        def_f, def_t = get_default_date_range()
        f_date = f_date_str or def_f
        t_date = t_date_str or def_t

        print(f"[*] Querying Stock Movement Report (222) Finished Goods (FAN Floor) from {f_date} to {t_date}...")
        payload = {
            "report": "222",
            "group_for": "3",  # FAN
            "item_sub_group": "100100000",  # Finished Goods
            "warehouse_id": "68",  # FAN Floor
            "f_date": f_date,
            "t_date": t_date,
            "submit": "Report"
        }

        res = self.session.post(WH_REPORT_URL, data=payload, timeout=60)
        soup = BeautifulSoup(res.text, "html.parser")
        tbl = soup.find("table")
        if not tbl:
            raise Exception("Stock Movement table not found in ERP response.")

        h2 = soup.find("h2")
        h3 = soup.find("h3")
        self.report_title = h2.get_text(" ", strip=True) if h2 else "Stock Movement Report"
        self.report_subtitle = h3.get_text(" ", strip=True) if h3 else f"Data Interval: {f_date} to {t_date}"

        rows = tbl.find_all("tr")
        parsed = []
        raw_rows = []

        # Row 0 & 1: Multilevel Header
        if len(rows) >= 2:
            h_row_0 = []
            for cell in rows[0].find_all(["th", "td"]):
                h_row_0.append({
                    "text": cell.get_text(" ", strip=True),
                    "bgcolor": cell.get("bgcolor", "") or "",
                    "colspan": int(cell.get("colspan", 1)),
                    "rowspan": int(cell.get("rowspan", 1)),
                    "style": cell.get("style", "") or ""
                })
            h_row_1 = []
            for cell in rows[1].find_all(["th", "td"]):
                h_row_1.append({
                    "text": cell.get_text(" ", strip=True),
                    "bgcolor": cell.get("bgcolor", "") or "",
                    "colspan": int(cell.get("colspan", 1)),
                    "rowspan": int(cell.get("rowspan", 1)),
                    "style": cell.get("style", "") or ""
                })
            self.raw_headers = [h_row_0, h_row_1]

        # Data starts from Row 2
        for r in rows[2:]:
            cols = [c.get_text(" ", strip=True) for c in r.find_all(["th", "td"])]
            if cols and cols[0].isdigit():
                raw_rows.append(cols)
                try:
                    opening = float(cols[7].replace(",", "") or 0)
                    local_purchase = float(cols[8].replace(",", "") or 0)
                    lc_purchase = float(cols[9].replace(",", "") or 0)
                    production = float(cols[10].replace(",", "") or 0)
                    sales_return = float(cols[11].replace(",", "") or 0)
                    transfer_prod_floor = float(cols[12].replace(",", "") or 0)
                    transfer_dist_point = float(cols[13].replace(",", "") or 0)
                    transfer_others = float(cols[14].replace(",", "") or 0)
                    transfer_total = float(cols[15].replace(",", "") or 0)
                    other_receive = float(cols[16].replace(",", "") or 0)
                    total_stock = float(cols[17].replace(",", "") or 0)
                    line_issue = float(cols[18].replace(",", "") or 0)
                    sales = float(cols[19].replace(",", "") or 0)
                    transfer_out = float(cols[20].replace(",", "") or 0)
                    other_issue = float(cols[21].replace(",", "") or 0)
                    out_total = float(cols[22].replace(",", "") or 0)
                    closing = float(cols[23].replace(",", "") or 0)
                    bin_closing = float(cols[24].replace(",", "") or 0)
                    diff = float(cols[25].replace(",", "") or 0)
                    rate = float(cols[26].replace(",", "") or 0)
                except (ValueError, IndexError):
                    continue

                parsed.append({
                    "sl": cols[0],
                    "company": cols[1],
                    "item_group": cols[2],
                    "product_category": cols[3],
                    "code": cols[4],
                    "item_name": cols[5],
                    "unit": cols[6] or "Pcs",
                    "opening": opening,
                    "local_purchase": local_purchase,
                    "lc_purchase": lc_purchase,
                    "production": production,
                    "sales_return": sales_return,
                    "transfer_prod_floor": transfer_prod_floor,
                    "transfer_dist_point": transfer_dist_point,
                    "transfer_others": transfer_others,
                    "transfer_total": transfer_total,
                    "other_receive": other_receive,
                    "total_stock": total_stock,
                    "line_issue": line_issue,
                    "sales": sales,
                    "transfer_out": transfer_out,
                    "other_issue": other_issue,
                    "out_total": out_total,
                    "closing": closing,
                    "bin_closing": bin_closing,
                    "diff": diff,
                    "rate": rate,
                    "remarks": cols[27] if len(cols) > 27 else ""
                })

        self.raw_rows = raw_rows

        # Capture total summary row from the bottom if present
        self.total_row = []
        if len(rows) > 2:
            last_cols = [c.get_text(" ", strip=True) for c in rows[-1].find_all(["th", "td"])]
            if any(last_cols) and not last_cols[0].isdigit():
                self.total_row = last_cols

        self.items = parsed
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        self.meta = {
            "title": "Stock Movement Report (222) — Finished Goods (FAN Floor)",
            "short_title": "Stock Movement (222) Finished Goods",
            "source": "Warehouse Module ⟶ Warehouse Reports ⟶ Stock Movement Report(222) ⟶ Item Sub Group: Finished Goods ⟶ Warehouse: FAN Floor",
            "company": "FAN",
            "report_code": "222",
            "item_sub_group": "Finished Goods",
            "warehouse_name": "FAN Floor",
            "f_date": f_date,
            "t_date": t_date,
            "collected_at": now_str,
            "total_items": len(parsed),
            "active_items": sum(1 for it in parsed if (it["opening"] > 0 or it["total_stock"] > 0 or it["out_total"] > 0 or it["closing"] > 0)),
            "total_opening": sum(it["opening"] for it in parsed),
            "total_production": sum(it["production"] for it in parsed),
            "total_transfer_in": sum(it["transfer_total"] for it in parsed),
            "total_stock": sum(it["total_stock"] for it in parsed),
            "total_out": sum(it["out_total"] for it in parsed),
            "total_closing": sum(it["closing"] for it in parsed)
        }

        print(f"[+] Collected {len(parsed)} items for Finished Goods FAN Floor (Closing Stock: {self.meta['total_closing']:,.2f})")
        return self.items

    def save_cache(self, filepath=STOCK_MOVEMENT_CACHE):
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump({
                "meta": self.meta,
                "items": self.items,
                "raw_headers": self.raw_headers,
                "raw_rows": self.raw_rows,
                "total_row": self.total_row,
                "report_title": self.report_title,
                "report_subtitle": self.report_subtitle
            }, f, indent=2, ensure_ascii=False)
        print(f"[+] Saved cache to {filepath}")

    def load_cache(self, filepath=STOCK_MOVEMENT_CACHE):
        if os.path.exists(filepath):
            with open(filepath, "r", encoding="utf-8") as f:
                d = json.load(f)
                self.items = d.get("items", [])
                self.meta = d.get("meta", {})
                self.raw_headers = d.get("raw_headers", [])
                self.raw_rows = d.get("raw_rows", [])
                self.total_row = d.get("total_row", [])
                self.report_title = d.get("report_title", "")
                self.report_subtitle = d.get("report_subtitle", "")
                return True
        return False

    def export_to_excel(self, output_file=STOCK_MOVEMENT_EXCEL):
        print(f"[*] Exporting Finished Goods Stock Movement Workbook to {output_file}...")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        # 1. Executive Summary Sheet
        ws_sum = wb.create_sheet(title="Executive Summary")
        ws_sum.views.sheetView[0].showGridLines = True
        self._build_summary_sheet(ws_sum)

        # 2. Master Stock Movement Table
        ws_main = wb.create_sheet(title="Stock Movement")
        ws_main.views.sheetView[0].showGridLines = True
        self._build_table_sheet(ws_main)

        wb.save(output_file)
        print(f"[+] Successfully generated {output_file}!")
        return output_file

    def _build_summary_sheet(self, ws):
        ws.merge_cells("A1:H2")
        b = ws["A1"]
        b.value = "MEP GROUP  |  FAN Finished Goods (FAN Floor) — Stock Movement Report (222)"
        b.font = Font(name=FONT_NAME, size=15, bold=True, color="FFFFFF")
        b.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        b.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        f_d = self.meta.get("f_date", "")
        t_d = self.meta.get("t_date", "")
        ws["A3"].value = f"{self.meta.get('source', '')}  |  Date Range: {f_d} to {t_d}"
        ws["A3"].font = Font(name=FONT_NAME, size=9.5, italic=True, color="64748B")

        kpis = [
            ("Total Catalog Items", len(self.items), "Finished Goods Catalog", "1B365D"),
            ("Active Movement Items", self.meta.get("active_items", 0), "Items with In/Out Activity", "059669"),
            ("Opening Quantity", f"{self.meta.get('total_opening', 0):,.2f}", "Month-Start Floor Inventory", "0284C7"),
            ("Total Closing Stock", f"{self.meta.get('total_closing', 0):,.2f}", "Net Available Finished Stock", "D97706"),
        ]

        col_pairs = [("B", "C"), ("D", "E"), ("F", "G"), ("H", "I")]
        for idx, (lbl, val, sub, col_hex) in enumerate(kpis):
            c1, c2 = col_pairs[idx]
            ws.merge_cells(f"{c1}5:{c2}5")
            ws.merge_cells(f"{c1}6:{c2}6")
            ws.merge_cells(f"{c1}7:{c2}7")

            v_cell = ws[f"{c1}5"]
            v_cell.value = val
            v_cell.font = Font(name=FONT_NAME, size=16, bold=True, color=col_hex)
            v_cell.alignment = Alignment(horizontal="center", vertical="center")

            l_cell = ws[f"{c1}6"]
            l_cell.value = lbl
            l_cell.font = Font(name=FONT_NAME, size=10, bold=True, color="334155")
            l_cell.alignment = Alignment(horizontal="center", vertical="center")

            s_cell = ws[f"{c1}7"]
            s_cell.value = sub
            s_cell.font = Font(name=FONT_NAME, size=8, color="64748B")
            s_cell.alignment = Alignment(horizontal="center", vertical="center")

            for r in range(5, 8):
                for c in [c1, c2]:
                    ws[f"{c}{r}"].fill = PatternFill(start_color=COLOR_LIGHT_BG, end_color=COLOR_LIGHT_BG, fill_type="solid")
                    ws[f"{c}{r}"].border = thin_border

        for col in ["B", "C", "D", "E", "F", "G", "H", "I"]:
            ws.column_dimensions[col].width = 18

    def _build_table_sheet(self, ws):
        ws.merge_cells("A1:Q1")
        title = ws["A1"]
        title.value = "MEP GROUP  |  FAN Finished Goods (FAN Floor) — Stock Movement Report (222)"
        title.font = Font(name=FONT_NAME, size=13, bold=True, color="FFFFFF")
        title.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 28

        headers = [
            ("SL", 7, "center"),
            ("Category", 16, "left"),
            ("Code", 18, "center"),
            ("Item Name", 36, "left"),
            ("Unit", 8, "center"),
            ("Opening", 14, "right"),
            ("Prod.", 12, "right"),
            ("Transfer In", 14, "right"),
            ("Other In", 12, "right"),
            ("Total Stock", 15, "right"),
            ("Line Issue", 13, "right"),
            ("Sales", 13, "right"),
            ("Transfer Out", 14, "right"),
            ("OUT Total", 15, "right"),
            ("Closing", 15, "right"),
            ("Rate (BDT)", 14, "right"),
            ("Remarks", 15, "left")
        ]

        ws.row_dimensions[3].height = 24
        for c_idx, (h_title, _, align) in enumerate(headers, start=1):
            c = ws.cell(row=3, column=c_idx, value=h_title)
            c.font = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
            c.alignment = Alignment(horizontal=align, vertical="center")
            c.border = header_border

        curr_r = 4
        for it in self.items:
            bg = COLOR_ZEBRA if curr_r % 2 == 1 else "FFFFFF"
            ws.row_dimensions[curr_r].height = 20
            row_vals = [
                it["sl"],
                it["product_category"],
                it["code"],
                it["item_name"],
                it["unit"],
                it["opening"],
                it["production"],
                it["transfer_total"],
                it["other_receive"],
                it["total_stock"],
                it["line_issue"],
                it["sales"],
                it["transfer_out"],
                it["out_total"],
                it["closing"],
                it["rate"],
                it["remarks"]
            ]
            for c_idx, v in enumerate(row_vals, start=1):
                c = ws.cell(row=curr_r, column=c_idx, value=v)
                c.font = Font(name=FONT_NAME, size=9)
                c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                c.border = thin_border
                c.alignment = Alignment(horizontal=headers[c_idx-1][2], vertical="center")
                if c_idx in [6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16]:
                    c.number_format = "#,##0.00"
            curr_r += 1

        # Summary Row
        ws.row_dimensions[curr_r].height = 24
        ws.cell(row=curr_r, column=1, value="Total")
        ws.cell(row=curr_r, column=4, value=f"{len(self.items)} Items")
        ws.cell(row=curr_r, column=6, value=self.meta.get("total_opening", 0))
        ws.cell(row=curr_r, column=10, value=self.meta.get("total_stock", 0))
        ws.cell(row=curr_r, column=14, value=self.meta.get("total_out", 0))
        ws.cell(row=curr_r, column=15, value=self.meta.get("total_closing", 0))

        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(row=curr_r, column=c_idx)
            c.font = Font(name=FONT_NAME, size=10, bold=True, color="000000")
            c.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            c.border = Border(top=Side(style='thin'), bottom=Side(style='double'))
            if c_idx in [6, 10, 14, 15]:
                c.number_format = "#,##0.00"

        for c_idx, (_, width, _) in enumerate(headers, start=1):
            col_letter = get_column_letter(c_idx)
            ws.column_dimensions[col_letter].width = width
